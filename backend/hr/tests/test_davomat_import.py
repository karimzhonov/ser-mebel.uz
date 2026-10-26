import json
from datetime import date

import pytest
from openpyxl import Workbook

from accounting.finance.importer import XlsxSource
from hr import services as S
from hr.importer import DavomatImporter
from hr.models import (
    Attendance,
    Employee,
    Holiday,
    HrSetting,
    Leave,
    PointAdjustment,
    Winner,
    WorkGroup,
)

pytestmark = pytest.mark.django_db

HEAD = [
    "Sana",
    "Hafta kuni",
    "Filial (guruh)",
    "Xodim",
    "Username",
    "Holat",
    "Kelgan vaqti",
    "Kechikish (daq)",
    "Kechikish sababi",
    "Izoh",
    "Aytgan vaqti",
    "Forma",
    "Telegram ID",
    "Ketgan vaqti",
    "Erta ketish (daq)",
    "Erta ketish sababi",
    "Ish vaqti",
    "Ball",
]


@pytest.fixture
def book(tmp_path):
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Davomat")
    ws.append(HEAD + ["", "Oylik reyting: oktabr 2026"])  # the summary block lives in column T…
    ws.append(
        [
            "2026-10-23",
            "Juma",
            "Zavod",
            "Jasur",
            "@jasur",
            "Vaqtida",
            "08:50",
            "",
            "",
            "",
            "",
            "Ha",
            "500",
            "18:02",
            "",
            "",
            "09:00–18:00",
            19,
        ]
    )
    ws.append(
        [
            "2026-10-23",
            "Juma",
            "Zavod",
            "Olim",
            "",
            "Kech qoldi",
            "09:40",
            40,
            "Sababsiz",
            "tirbandlik",
            "",
            "",
            "501",
            "",
            "",
            "",
            "09:00–18:00",
            -7,
        ]
    )
    ws.append(
        [
            "2026-10-23",
            "Juma",
            "Zavod",
            "Bek",
            "",
            "Kelmadi",
            "",
            "",
            "",
            "",
            "",
            "",
            "502",
            "",
            "",
            "",
            "09:00–18:00",
            -15,
        ]
    )
    ws.append(
        [
            "2026-10-23",
            "Juma",
            "Zavod",
            "Dilnoza",
            "",
            "Sababli kelmadi",
            "",
            "",
            "",
            "Kasal",
            "",
            "",
            "503",
            "",
            "",
            "",
            "09:00–18:00",
            "",
        ]
    )
    ws.append(
        [
            "2026-10-24",
            "Shanba",
            "Zavod",
            "Jasur",
            "@jasur",
            "Dam kuni",
            "",
            "",
            "",
            "",
            "",
            "",
            "500",
            "",
            "",
            "",
            "09:00–18:00",
            "",
        ]
    )
    ws.append(
        [
            "2026-10-24",
            "Shanba",
            "Zavod",
            "Olim",
            "",
            "Vaqtida",
            "08:59",
            "",
            "",
            "",
            "",
            "",
            "501",
            "16:00",
            120,
            "Sababli: shifokor",
            "09:00–18:00",
            10,
        ]
    )
    ws.append(
        ["", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", 1, "Jasur"]
    )  # summary-only row
    ws.append(
        [
            "2026-10-25",
            "Yakshanba",
            "Zavod",
            "Xato",
            "",
            "Nimadir",
            "",
            "",
            "",
            "",
            "",
            "",
            "509",
            "",
            "",
            "",
            "",
            "",
        ]
    )
    ws.append(
        [
            "2026-10-22",
            "Payshanba",
            "Zavod",
            "Jasur",
            "@jasur",
            "Vaqtida",
            "08:58",
            "",
            "",
            "",
            "",
            "",
            "500",
            "18:00",
            "",
            "",
            "09:00–18:00",
            99,
        ]
    )  # wrong Ball
    ws.append(
        [
            "2026-10-26",
            "Dushanba",
            "Zavod",
            "Olim",
            "",
            "Hali kelmadi",
            "",
            "",
            "",
            "",
            "",
            "",
            "501",
            "",
            "",
            "",
            "09:00–18:00",
            "",
        ]
    )  # the open day
    ex = wb.create_sheet("DavomatExport")
    ex.append(["Kalit", "Qiymat"])
    props = {
        "dv_groups": json.dumps({"-100777": "Zavod"}),
        "dv_e_500": json.dumps(
            {
                "tg": 500,
                "id": 1,
                "n": "Jasur Karimov",
                "u": "jasur",
                "a": 1,
                "dm": 1,
                "off": [7],
                "g": -100777,
                "c": "2026-09-01",
                "ws": "08:30",
                "we": "17:30",
            }
        ),
        "dv_e_501": json.dumps(
            {
                "tg": 501,
                "id": 2,
                "n": "Olim",
                "u": "",
                "a": 1,
                "dm": 0,
                "off": None,
                "g": -100777,
                "c": "2026-09-01",
            }
        ),
        "dv_e_502": json.dumps({"tg": 502, "id": 3, "n": "Bek", "a": 0, "g": None}),
        "dv_s_GRACE_MIN": "10",
        "dv_holidays": json.dumps({"2027-01-01": "Yangi yil"}),
        "dv_hd_2027-01-01": "1",
        "dv_leaves": json.dumps(
            [{"id": 4, "tg": 501, "from": "2026-11-02", "to": "2026-11-05", "r": "ta'til"}]
        ),
        "dv_adj_2026-10": json.dumps(
            [{"tg": 500, "d": "2026-10-20", "v": 10, "r": "Mijozdan rahmat"}]
        ),
        "dv_win_m_2026-09": "500",
        "dv_a_2026-10-26_500": json.dumps(
            {"in": "08:55", "st": "on_time", "late": 0, "ws": "08:30", "we": "17:30", "uni": 1}
        ),
        "dv_curday": "2026-10-26",
        "dv_since": "2026-09-01",
    }
    for k, v in props.items():
        ex.append([k, v])
    p = tmp_path / "davomat.xlsx"
    wb.save(p)
    return str(p)


def test_import_davomat(book):
    rep = DavomatImporter(XlsxSource(book)).run()
    assert rep.ok(), rep.text()
    jasur = Employee.objects.get(telegram_id=500)
    assert (jasur.name, jasur.number, jasur.off_days, jasur.work_start, jasur.can_dm) == (
        "Jasur Karimov",
        1,
        [7],
        "08:30",
        None,
    )
    assert jasur.group == WorkGroup.objects.get(chat_id=-100777)
    assert not Employee.objects.get(telegram_id=502).is_active
    assert not Employee.objects.get(telegram_id=503).is_active  # rebuilt from the sheet
    assert Attendance.objects.filter(finalized=True).count() == 7
    wrong = Attendance.objects.get(employee__telegram_id=500, day=date(2026, 10, 22))
    assert wrong.points == 15 and any("Ball" in w for w in rep.warnings)  # recomputed, not copied
    live = Attendance.objects.get(employee__telegram_id=501, day=date(2026, 10, 26))
    assert not live.finalized and live.status == ""  # the open day stays live
    olim = Attendance.objects.get(employee__telegram_id=501, day=date(2026, 10, 23))
    assert (olim.status, olim.late_min, olim.late_reason, olim.late_comment, olim.points) == (
        "late",
        40,
        "Sababsiz",
        "tirbandlik",
        -7,
    )
    early = Attendance.objects.get(employee__telegram_id=501, day=date(2026, 10, 24))
    assert (early.early_min, early.early_reason, early.early_comment) == (
        120,
        "Sababli",
        "shifokor",
    )
    assert Attendance.objects.get(employee__telegram_id=503).excuse == "Kasal"
    today = Attendance.objects.get(employee__telegram_id=500, day=date(2026, 10, 26))
    assert not today.finalized and today.status == "on_time" and today.uniform is True
    assert any("Nimadir" in w for w in rep.warnings)
    assert HrSetting.objects.get(key="GRACE_MIN").value == "10"
    assert Holiday.objects.filter(day=date(2027, 1, 1)).exists()
    assert Leave.objects.get(number=4).employee.telegram_id == 501
    assert PointAdjustment.objects.get().value == 10
    assert Winner.objects.get(kind="m", period="2026-09").employee == jasur
    assert S.since() == date(2026, 9, 1) and S.state("CURDAY") == S.now().date().isoformat()


def test_reimport_is_idempotent(book):
    DavomatImporter(XlsxSource(book)).run()
    counts = (Employee.objects.count(), Attendance.objects.count(), PointAdjustment.objects.count())
    rep = DavomatImporter(XlsxSource(book)).run()
    assert rep.ok() and counts == (
        Employee.objects.count(),
        Attendance.objects.count(),
        PointAdjustment.objects.count(),
    )


def test_import_without_export_sheet_warns(tmp_path):
    wb = Workbook()
    wb.active.title = "Davomat"
    wb.active.append(HEAD)
    p = tmp_path / "x.xlsx"
    wb.save(p)
    rep = DavomatImporter(XlsxSource(str(p))).run()
    assert any("DavomatExport" in w for w in rep.warnings)


def test_finance_import_never_reads_davomat(book, db):
    """The two importers are separate: the finance one ignores these sheets."""
    from accounting.finance.importer import Importer
    from accounting.finance_models import Operation

    Importer(XlsxSource(book)).run()
    assert not Operation.objects.exists() and not Attendance.objects.exists()


def test_reimport_keeps_what_the_new_bot_wrote(book):
    """Go-live day: the new bot already recorded a check-in; a second import must not
    replace it with the Apps Script's stale memory row."""
    DavomatImporter(XlsxSource(book)).run()
    a = Attendance.objects.get(employee__telegram_id=500, day=date(2026, 10, 26))
    a.status, a.late_min = "late", 25
    a.save()
    DavomatImporter(XlsxSource(book)).run()
    a.refresh_from_db()
    assert (a.status, a.late_min) == ("late", 25)

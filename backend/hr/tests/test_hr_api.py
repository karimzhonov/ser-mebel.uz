from datetime import date

import pytest

from hr import services as S
from hr.models import Attendance, Employee, Leave, PointAdjustment, Winner

from .conftest import GROUP

pytestmark = pytest.mark.django_db


@pytest.fixture
def admin(api):
    return api(superuser=True)


@pytest.fixture
def jasur(group):
    e, _ = S.upsert_emp({"id": 500, "first_name": "Jasur"}, group)
    return e


def test_requires_permission(api, jasur):
    assert api().get("/api/v1/hr/day/").status_code == 403
    assert api("hr.view_attendance").get("/api/v1/hr/day/").status_code == 200


def test_day_board(admin, jasur, clock):
    clock.set(12, 0)
    S.set_att(
        clock.t.date(),
        jasur,
        status="late",
        in_time="09:30",
        late_min=30,
        work_start="09:00",
        work_end="18:00",
    )
    r = admin.get("/api/v1/hr/day/").data
    row = r["rows"][0]
    assert row["status"] == "late" and row["in_time"] == "09:30" and row["points"] == -5
    assert r["counts"] == {"late": 1}


def test_edit_finalized_day_recomputes_points(admin, jasur):
    a = Attendance.objects.create(
        employee=jasur,
        day=date(2026, 10, 20),
        status="absent",
        points=-15,
        finalized=True,
        work_start="09:00",
        work_end="18:00",
    )
    r = admin.patch(
        f"/api/v1/hr/attendance/{a.id}/", {"status": "excused", "excuse": "Kasal"}, format="json"
    )
    assert r.status_code == 200 and r.data["points"] is None
    r = admin.patch(
        f"/api/v1/hr/attendance/{a.id}/",
        {"status": "on_time", "in_time": "08:55", "out_time": "18:00"},
        format="json",
    )
    assert r.data["points"] == 15
    assert admin.patch(
        f"/api/v1/hr/attendance/{a.id}/", {"day": "2026-01-01"}, format="json"
    ).status_code in (200, 400)
    assert Attendance.objects.get(pk=a.id).day == date(2026, 10, 20)


def test_employee_edit_hours_and_off_days(admin, jasur):
    url = f"/api/v1/hr/employees/{jasur.id}/"
    assert (
        admin.patch(url, {"work_start": "10:00", "work_end": "09:00"}, format="json").status_code
        == 400
    )
    assert admin.patch(url, {"off_days": [8]}, format="json").status_code == 400
    r = admin.patch(
        url, {"work_start": "10:00", "work_end": "19:00", "off_days": [7, 6]}, format="json"
    )
    assert (
        r.status_code == 200 and r.data["hours"] == "10:00–19:00" and r.data["off_days"] == [6, 7]
    )


def test_leaves_points_winners(admin, jasur, bot_events):
    r = admin.post(
        "/api/v1/hr/leaves/",
        {
            "employee": jasur.id,
            "date_from": "2026-11-02",
            "date_to": "2026-11-04",
            "reason": "ta'til",
        },
        format="json",
    )
    assert r.status_code == 201 and Leave.objects.get().number == 1
    r = admin.post(
        "/api/v1/hr/points/",
        {"employee": jasur.id, "day": "2026-10-26", "value": 10, "reason": "rahmat"},
        format="json",
    )
    assert r.status_code == 201 and PointAdjustment.objects.count() == 1
    pe = bot_events.one("hr.points_added")
    assert pe["from_site"] is True and pe["value"] == 10 and pe["employee"]["telegram_id"] == 500
    assert (
        admin.post(
            "/api/v1/hr/points/",
            {"employee": jasur.id, "day": "2026-10-26", "value": 500},
            format="json",
        ).status_code
        == 400
    )
    r = admin.post(
        "/api/v1/hr/winners/announce/",
        {"kind": "m", "period": "2026-10", "employee": jasur.id},
        format="json",
    )
    assert r.status_code == 201 and Winner.objects.count() == 1
    won = bot_events.one("hr.winner_announced")
    assert won["groups"] == [GROUP] and won["kind"] == "m" and won["employee"]["telegram_id"] == 500
    assert (
        admin.post(
            "/api/v1/hr/winners/announce/",
            {"kind": "m", "period": "2026-10", "employee": jasur.id},
            format="json",
        ).status_code
        == 400
    )


def test_stats_settings_export(admin, jasur, clock):  # 08:00 — today still «waiting», 0 points
    Attendance.objects.create(
        employee=jasur, day=date(2026, 10, 20), status="on_time", points=17, finalized=True
    )
    r = admin.get("/api/v1/hr/stats/?month=2026-10").data
    assert r["rows"][0]["points"] == 17
    assert admin.get("/api/v1/hr/stats/?year=2026").status_code == 200
    assert (
        admin.patch("/api/v1/hr/settings/", {"GRACE_MIN": "abc"}, format="json").status_code == 400
    )
    r = admin.patch(
        "/api/v1/hr/settings/", {"GRACE_MIN": "10", "WORK_START": "9:00"}, format="json"
    )
    assert r.data["values"]["GRACE_MIN"] == "10" and r.data["values"]["WORK_START"] == "09:00"
    x = admin.get("/api/v1/hr/export/?date_from=2026-10-01&date_to=2026-10-31")
    assert x.status_code == 200 and x.content[:2] == b"PK"


def test_hr_data_never_mixes_with_finance(admin, jasur):
    """Davomat employees are not finance wallets and vice versa."""
    from accounting import services as F

    assert (
        jasur.member is None and F.wallets() == [] or all(m.telegram_id != 500 for m in F.wallets())
    )
    assert not Employee.objects.filter(telegram_id=100).exists()  # the boss is not an employee


def test_ensure_row_for_absent_day_then_excuse(admin, jasur):
    jasur.created_on = date(2026, 10, 1)
    jasur.save()
    r = admin.post(
        "/api/v1/hr/attendance/ensure/", {"employee": jasur.id, "day": "2026-10-20"}, format="json"
    )
    assert r.status_code == 200 and r.data["finalized"] and r.data["status"] in ("absent", "off")
    r2 = admin.patch(
        f"/api/v1/hr/attendance/{r.data['id']}/",
        {"status": "excused", "excuse": "Kasal"},
        format="json",
    )
    assert r2.data["points"] is None


def test_ensure_refuses_days_before_joining(admin, jasur):
    jasur.created_on = date(2026, 10, 1)
    jasur.save()
    r = admin.post(
        "/api/v1/hr/attendance/ensure/", {"employee": jasur.id, "day": "2026-09-15"}, format="json"
    )
    assert r.status_code == 400 and not Attendance.objects.exists()


def test_new_arrival_time_decides_late(admin, jasur):
    a = Attendance.objects.create(
        employee=jasur,
        day=date(2026, 10, 20),
        status="absent",
        finalized=True,
        work_start="09:00",
        work_end="18:00",
    )
    r = admin.patch(
        f"/api/v1/hr/attendance/{a.id}/", {"in_time": "09:40", "out_time": "18:00"}, format="json"
    )
    assert r.data["status"] == "late" and r.data["late_min"] == 40 and r.data["points"] == -5 + 5
    assert (
        admin.patch(
            f"/api/v1/hr/attendance/{a.id}/", {"status": "waiting"}, format="json"
        ).status_code
        == 400
    )


def test_excel_names_cannot_become_formulas(admin, jasur, clock):
    from io import BytesIO

    from openpyxl import load_workbook

    jasur.name = '=HYPERLINK("http://evil","x")'
    jasur.save()
    Attendance.objects.create(
        employee=jasur, day=date(2026, 10, 20), status="on_time", finalized=True, points=10
    )
    x = admin.get("/api/v1/hr/export/?date_from=2026-10-01&date_to=2026-10-25")
    ws = load_workbook(BytesIO(x.content))["Davomat"]
    assert str(ws["D2"].value).startswith("'=")


def test_identity_fields_are_read_only(admin, jasur):
    r = admin.patch(f"/api/v1/hr/employees/{jasur.id}/", {"telegram_id": 12345}, format="json")
    assert r.status_code == 400

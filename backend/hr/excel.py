"""Excel of the davomat — same columns as the old «Davomat» sheet + the ranking."""

from __future__ import annotations

import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from accounting.finance.text import safe_cell as T

from . import services as S
from .constants import ST_PLAIN, WEEKDAYS
from .models import Attendance

HEAD = [
    "Sana",
    "Hafta kuni",
    "Filial (guruh)",
    "Xodim",
    "Username",
    "Holat",
    "Kelgan vaqti",
    "Kechikish (daq)",
    "Kechikish sababi",
    "Izoh",
    "Aytgan vaqti",
    "Forma",
    "Telegram ID",
    "Ketgan vaqti",
    "Erta ketish (daq)",
    "Erta ketish sababi",
    "Ish vaqti",
    "Ball",
]
RANK = [
    "O'rin",
    "Xodim",
    "Ball",
    "Ish kunlari",
    "Vaqtida",
    "Kech qoldi",
    "Jami kechikish (daq)",
    "Kelmadi",
    "Sababli kelmadi",
    "Erta ketdi",
    "Intizom %",
    "Qo'shimcha ball",
]


def _head(ws, values):
    ws.append(values)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
    ws.freeze_panes = "A2"


def build_xlsx(start: date, end: date) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Davomat"
    _head(ws, HEAD)
    now_d = S.now()
    rows = list(
        Attendance.objects.filter(day__gte=start, day__lte=min(end, now_d.date()), finalized=True)
        .select_related("employee")
        .order_by("day", "employee__name")
    )
    if start <= now_d.date() <= end:
        live = S.day_rows(now_d.date(), now_d)
    else:
        live = []
    for a in rows:
        ws.append(
            [
                a.day,
                WEEKDAYS[a.day.isoweekday() - 1],
                T(a.group_title),
                T(a.employee.name),
                f"@{a.employee.username}" if a.employee.username else "",
                ST_PLAIN.get(a.status, a.status),
                S.hm(a.in_time),
                a.late_min if a.status == "late" else "",
                a.late_reason if a.status == "late" else "",
                T(a.late_comment if a.status == "late" else (a.note or a.excuse)),
                S.hm(a.eta) if not a.in_time else "",
                "Ha" if a.uniform is True else ("Yo'q" if a.uniform is False else ""),
                str(a.employee.telegram_id),
                S.hm(a.out_time),
                a.early_min or "",
                (
                    T(
                        (a.early_reason or "Sababsiz")
                        + (f": {a.early_comment}" if a.early_comment else "")
                    )
                    if a.early_min
                    else ""
                ),
                f"{a.work_start}–{a.work_end}" if a.work_start else "",
                a.points if a.points is not None else "",
            ]
        )
    today = now_d.date()
    for r in live:
        ws.append(
            [
                today,
                WEEKDAYS[today.isoweekday() - 1],
                T(r["group"]),
                T(r["name"]),
                f"@{r['username']}" if r["username"] else "",
                ST_PLAIN.get(r["status"], r["status"]),
                r["in_time"],
                r["late_min"] or "",
                r["late_reason"],
                T(r["late_comment"] or r["note"]),
                r["eta"],
                "Ha" if r["uniform"] is True else ("Yo'q" if r["uniform"] is False else ""),
                "",
                r["out_time"],
                r["early_min"] or "",
                r["early_reason"],
                r["hours"],
                r["points"] if r["points"] is not None else "",
            ]
        )
    for i, w in enumerate(
        [12, 12, 16, 24, 14, 18, 10, 10, 12, 30, 10, 8, 14, 10, 10, 24, 12, 8], 1
    ):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2, max_col=1):
        row[0].number_format = "dd.mm.yyyy"

    rk = wb.create_sheet("Reyting")
    _head(rk, RANK)
    for i, s in enumerate(S.period_stats(start, end, now_d), 1):
        rk.append(
            [
                i,
                T(s.e.name),
                s.pts,
                s.exp,
                s.on_time,
                s.late,
                s.late_min,
                s.absent,
                s.excused,
                s.early_n,
                s.disc,
                s.adj + s.bonus,
            ]
        )
    rk.column_dimensions["B"].width = 26
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

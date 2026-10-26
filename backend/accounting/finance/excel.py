"""«📥 Excel»: the month report file (what the «Xulosa» and «Tekshiruv» sheets showed)."""

from __future__ import annotations

import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .. import services as S
from ..finance_constants import CAT_SALARY, OpKind
from ..finance_models import MonthReport
from .text import add_months, month_bounds, month_full, month_name, safe_cell

DARK, MINT, BAND = "0B3D2E", "E8F3EC", "F3F8F5"
HEAD_FONT = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor=DARK)
MONEY = "#,##0;[Red]-#,##0"


def _head(ws, row, values):
    for i, v in enumerate(values, 1):
        c = ws.cell(row=row, column=i, value=v)
        c.font, c.fill = HEAD_FONT, HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center")


def _section(ws, row, text):
    ws.cell(row=row, column=1, value=text).font = Font(bold=True, color=DARK, size=12)


def _widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build_xlsx(start: date, end: date, title: str = "") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Xulosa"
    sm = S.summary(start, end)
    ws["A1"] = f"📊 SerMebel — Xulosa {title}".strip()
    ws["A1"].font = Font(bold=True, size=16, color=DARK)
    ws["A2"] = f"Davr: {start:%d.%m.%Y} — {end:%d.%m.%Y}"
    cards = [
        ("Xarajatlar", -sm["expense"]),
        ("Daromadlar", sm["income"]),
        ("Balans", sm["balance"]),
        ("Xodimlarga berildi", sm["given_to_staff"]),
        ("Rahbarga (shaxsiy)", sm["boss_personal"]),
        ("Uy uchun (shaxsiy)", sm["home"]),
    ]
    for i, (lab, val) in enumerate(cards):
        ws.cell(row=4, column=1 + i * 2, value=lab).font = Font(bold=True, color=DARK)
        c = ws.cell(row=5, column=1 + i * 2, value=float(val))
        c.number_format, c.font = MONEY, Font(bold=True, size=13)
    r = 7
    _section(ws, r, "Xarajatlar turlari bo'yicha")
    _head(ws, r + 1, ["Kategoriya", "Summa", "Operatsiya", "Ulushi"])
    r += 2
    tot = sum(x["sum"] for x in sm["expense_by_category"]) or 1
    for x in sm["expense_by_category"]:
        ws.append([f"{x['emoji']} {x['name']}", float(x["sum"]), x["n"], float(x["sum"] / tot)])
        ws.cell(row=r, column=2).number_format = MONEY
        ws.cell(row=r, column=4).number_format = "0%"
        r += 1
    r += 1
    _section(ws, r, "Daromadlar turlari bo'yicha")
    _head(ws, r + 1, ["Turi", "Summa", "Operatsiya"])
    r += 2
    for x in sm["income_by_category"]:
        ws.append([f"{x['emoji']} {x['name']}", float(x["sum"]), x["n"]])
        ws.cell(row=r, column=2).number_format = MONEY
        r += 1
    r += 1
    _section(ws, r, "Hamyonlar (xodimlar)")
    _head(
        ws, r + 1, ["Xodim", "Davrda berildi", "Davrda sarfladi", "Hamyonda (jami)", "Oylik hisob"]
    )
    r += 2
    bal = S.balances()
    ym = start.strftime("%Y-%m")
    reported = set(MonthReport.objects.filter(month=ym).values_list("member_id", flat=True))
    for m in S.wallets():
        ops = list(S.period_ops(start, end, m))
        given = sum(float(t.amount) for t in ops if t.kind == OpKind.INCOME)
        spent = sum(
            float(t.amount)
            for t in ops
            if t.kind == OpKind.EXPENSE and t.category.code != "otkazma"
        )
        ws.append(
            [
                ("💼 " if m.is_admin else "👤 ") + m.name,
                given,
                spent,
                "— kassa" if m.is_admin else float(bal.get(m.id, 0)),
                "—" if m.is_admin else ("✅ topshirdi" if m.id in reported else "⏳ yo'q"),
            ]
        )
        for col in (2, 3, 4):
            ws.cell(row=r, column=col).number_format = MONEY
        r += 1
    r += 1
    _section(ws, r, "👷 Oylik — kimga")
    _head(ws, r + 1, ["Ism", "Davrda berildi"])
    r += 2
    for name, val in sm["salary_by_worker"]:
        ws.append([name, float(val)])
        ws.cell(row=r, column=2).number_format = MONEY
        r += 1
    r += 1
    _section(ws, r, "Oylar bo'yicha (12 oy)")
    _head(ws, r + 1, ["Oy", "Daromad", "Xarajat", "Balans"])
    r += 2
    for k in range(-11, 1):
        mm = add_months(ym, k)
        s, e = month_bounds(mm)
        x = S.summary(s, e)
        ws.append([month_name(mm), float(x["income"]), float(x["expense"]), float(x["balance"])])
        for col in (2, 3, 4):
            ws.cell(row=r, column=col).number_format = MONEY
        r += 1
    _widths(ws, [34, 18, 16, 18, 18, 16, 16, 16, 16, 16, 16, 16])

    ws2 = wb.create_sheet("Amallar")
    _head(
        ws2,
        1,
        [
            "№",
            "Sana",
            "Vaqt",
            "Xodim",
            "Turi",
            "Summa (so'm)",
            "Valyuta",
            "Asl summa",
            "Kurs",
            "Kategoriya",
            "Buyurtma ID",
            "Izoh",
            "Ishchi / kimga",
            "Ish haqi turi",
            "Manba",
        ],
    )
    for i, t in enumerate(S.period_ops(start, end).order_by("date", "number"), 2):
        ws2.append(
            [
                t.number,
                t.date,
                t.time.strftime("%H:%M") if t.time else "",
                t.member.name,
                t.get_kind_display(),
                float(t.amount),
                t.currency,
                float(t.orig_amount) if t.orig_amount else None,
                float(t.rate) if t.rate else None,
                t.category.name,
                t.order_ref,
                safe_cell(t.comment),
                safe_cell(t.worker),
                t.pay_type if t.category.code == CAT_SALARY else "",
                t.get_source_display(),
            ]
        )
        ws2.cell(row=i, column=2).number_format = "dd.mm.yyyy"
        ws2.cell(row=i, column=6).number_format = MONEY
    ws2.freeze_panes = "A2"
    _widths(ws2, [8, 12, 8, 22, 10, 16, 8, 12, 10, 26, 12, 40, 24, 12, 12])

    ws3 = wb.create_sheet("Tekshiruv")
    _head(ws3, 1, ["Daraja", "Muammo", "Tafsilot", "Yechim"])
    lv = {"h": "🔴", "m": "🟠", "l": "🟡"}
    issues = S.checks()
    for i in issues:
        ws3.append([lv[i["level"]], i["title"], i["detail"], i["fix"]])
    if not issues:
        ws3.append(["", "✅ Kamchilik topilmadi"])
    for row in ws3.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    _widths(ws3, [8, 48, 52, 60])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def month_xlsx(ym: str) -> tuple[bytes, str]:
    s, e = month_bounds(ym)
    return build_xlsx(s, e, month_full(ym)), f"SerMebel_{ym}.xlsx"

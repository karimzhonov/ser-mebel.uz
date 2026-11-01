"""«Oylik chiqarish» on «👷 Ishchi kartasi»: every worker has their own Excel template
(fixed, bonus and task pay look different). The button fills in what the system knows
— the name and the advances paid since the last pay day — and the accountant fills in
the rest by hand. The day before a worker's pay day the accountant (TelegramMember
`is_accountant`, the admins when nobody is) gets a bot reminder.

Placeholders a template may use (in any cell, `{x}` or `{{x}}`):
  {ism} {lavozim} {bolim} {oy} {sana} {davr} {davr_boshi} {davr_oxiri}
  {avans} (sum of advances) {avans_soni}
and, on one row, a list of the advances — the row is repeated downwards for each one:
  {avans_n} {avans_sana} {avans_summa} {avans_kim} {avans_izoh}
Templates without placeholders (like the ones the office uses: ODDIY, OFIS, MENEJER,
ZAMER, SBORKA) are read by their text: the value cell right of «Xodim (F.I.Sh.)»,
«Lavozim», «Soha / bo'lim», «Davr (oy 1-sanasi)», «Davr», «Hisob sanasi» is filled
when empty, «1 USD kursi» gets today's rate, and the numbered rows under the
«… AVANSLAR» title (columns Sana / Izoh / Kim oldi / Nomi / Valyuta / Summa / Kurs) get
the advances. Formulas, borders and drop-downs of the template are left as they are.
Whatever the template looks like, an «Avanslar» sheet with the list is added too.
"""

from __future__ import annotations

import io
import re
from copy import copy
from dataclasses import dataclass
from datetime import date, time, timedelta
from decimal import Decimal
from types import SimpleNamespace

from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException

from core.exceptions import ServiceError

from . import events as E
from . import services as S
from .constants import PayType
from .models import Worker
from .text import add_months, month_bounds, month_full, safe_cell, today

MAX_TEMPLATE = 5 * 1024 * 1024
XLSX_MAGIC = b"PK\x03\x04"
ACCOUNTANT = "is_accountant"


# ------------------------------------------------------------------ pay dates
def pay_date(worker: Worker, ym: str) -> date:
    """The pay day of `ym` (the 31st is the month's last day in short months);
    without a pay day — the month's last day."""
    _, end = month_bounds(ym)
    if not worker.pay_day:
        return end
    return end.replace(day=min(int(worker.pay_day), end.day))


def period(worker: Worker, ym: str) -> tuple[date, date]:
    """Advances that count for the salary of `ym`: from the day after the previous pay
    day up to this one. Without a pay day — the calendar month."""
    if not worker.pay_day:
        return month_bounds(ym)
    return pay_date(worker, add_months(ym, -1)) + timedelta(days=1), pay_date(worker, ym)


def next_month(worker: Worker, on: date | None = None) -> str:
    """The month to issue now. With a pay day: the one whose pay day is next (today
    included). Without one the salary is for a calendar month: the one just ended,
    or the current one from the 25th on."""
    t = on or today()
    ym = t.strftime("%Y-%m")
    if not worker.pay_day:
        return ym if t.day >= 25 else add_months(ym, -1)
    return ym if pay_date(worker, ym) >= t else add_months(ym, 1)


def _payments(worker: Worker, start: date, end: date) -> list:
    key = S._worker_key(worker.name)
    rows = [t for t in S.salary_ops(start, end) if S._worker_key(t.worker) == key]
    rows.sort(key=lambda t: (t.date, t.time or time.min, t.id))
    return rows


def advances(worker: Worker, start: date, end: date) -> list:
    return [t for t in _payments(worker, start, end) if t.pay_type == PayType.ADVANCE]


def untyped(worker: Worker, start: date, end: date) -> list:
    """Payments booked without «Avans / Oylik» (old rows, imports): not put into the
    sheet, but shown so the accountant can check them."""
    return [
        t
        for t in _payments(worker, start, end)
        if t.pay_type not in (PayType.ADVANCE, PayType.SALARY)
    ]


MAX_PERIOD_DAYS = 400


@dataclass
class Sheet:
    worker: Worker
    ym: str
    pay_date: date
    start: date
    end: date
    rows: list

    @property
    def total(self) -> Decimal:
        return sum((t.amount for t in self.rows), S.ZERO)


def salary_info(
    worker: Worker, ym: str | None = None, start: date | None = None, end: date | None = None
) -> Sheet:
    """`start`/`end` (both) replace the period worked out from the pay day."""
    ym = ym or next_month(worker)
    if start and end:
        if start > end:
            raise ServiceError("Davr boshi oxiridan keyin bo'lmasin.")
        if (end - start).days > MAX_PERIOD_DAYS:
            raise ServiceError("Davr juda uzun (ko'pi bilan bir yil).")
    else:
        start, end = period(worker, ym)
    return Sheet(worker, ym, pay_date(worker, ym), start, end, advances(worker, start, end))


# ------------------------------------------------------------------ the template
def save_template(worker: Worker, name: str, content: bytes) -> Worker:
    name = (name or "").strip() or "shablon.xlsx"
    if not name.lower().endswith(".xlsx"):
        raise ServiceError("Faqat .xlsx fayl (Excel → «Saqlash» → Excel kitobi .xlsx).")
    if not content:
        raise ServiceError("Fayl bo'sh.")
    if len(content) > MAX_TEMPLATE:
        raise ServiceError("Fayl juda katta (5 MB gacha).")
    if not content.startswith(XLSX_MAGIC):
        raise ServiceError("Bu Excel (.xlsx) fayl emas.")
    try:
        load_workbook(io.BytesIO(content))
    except (InvalidFileException, KeyError, ValueError, OSError) as e:
        raise ServiceError("Excel faylni ochib bo'lmadi — .xlsx ekanini tekshiring.") from e
    worker.salary_template = content
    worker.salary_template_name = name[-255:]
    worker.salary_template_at = timezone.now()
    worker.save(update_fields=["salary_template", "salary_template_name", "salary_template_at"])
    return worker


def remove_template(worker: Worker) -> Worker:
    worker.salary_template = None
    worker.salary_template_name = ""
    worker.salary_template_at = None
    worker.save(update_fields=["salary_template", "salary_template_name", "salary_template_at"])
    return worker


def template_file(worker: Worker) -> tuple[bytes, str]:
    """The uploaded template, or a starter one that shows the placeholders."""
    if worker.salary_template:
        return bytes(worker.salary_template), worker.salary_template_name or "shablon.xlsx"
    return _starter(), f"Oylik_shablon_{_slug(worker.name)}.xlsx"


def _slug(s: str) -> str:
    return re.sub(r"[^\w-]+", "_", s, flags=re.U).strip("_")[:60] or "ishchi"


def _starter() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Oylik"
    bold = Font(bold=True)
    head = PatternFill("solid", fgColor="E8F5E9")
    ws["A1"], ws["A1"].font = "Oylik varaqasi — {oy}", Font(bold=True, size=14)
    for r, (k, v) in enumerate(
        [
            ("Ism", "{ism}"),
            ("Lavozim", "{lavozim}"),
            ("Bo'lim", "{bolim}"),
            ("Oylik sanasi", "{sana}"),
            ("Davr", "{davr}"),
        ],
        start=3,
    ):
        ws.cell(r, 1, k).font = bold
        ws.cell(r, 2, v)
    for c, h in enumerate(["№", "Sana", "Avans (so'm)", "Kim berdi", "Izoh"], start=1):
        x = ws.cell(9, c, h)
        x.font, x.fill = bold, head
    for c, v in enumerate(
        ["{avans_n}", "{avans_sana}", "{avans_summa}", "{avans_kim}", "{avans_izoh}"], start=1
    ):
        ws.cell(10, c, v)
    rows = [
        ("Avanslar jami", "{avans}"),
        ("Oklad / fiks", ""),
        ("Bonus", ""),
        ("Vazifalar (ishbay)", ""),
        ("Ushlab qolish", ""),
        ("Qo'lga beriladi", ""),
    ]
    for r, (k, v) in enumerate(rows, start=13):
        ws.cell(r, 1, k).font = bold
        ws.cell(r, 3, v)
    for col, w in zip("ABCDE", (22, 14, 18, 22, 40)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ filling it in
PH = re.compile(r"\{\{?\s*([a-z_]+)\s*\}?\}")
ROW_KEYS = {
    "avans_n",
    "avans_sana",
    "avans_summa",
    "avans_kim",
    "avans_izoh",
    "avans_valyuta",
    "avans_asl",
    "avans_kurs",
}


def work_month(sh: Sheet) -> date:
    """The month the salary is for: the one holding most of the advance period
    (pay day 5 Nov → October, pay day 25 Nov → November)."""
    mid = sh.start + (sh.end - sh.start) / 2
    return mid.replace(day=1)


def _values(sh: Sheet) -> dict:
    w = sh.worker
    wm = work_month(sh)
    return {
        "ism": w.name,
        "lavozim": w.position,
        "bolim": w.group,
        "oy": month_full(wm.strftime("%Y-%m")),
        "oy_boshi": wm,
        "sana": sh.pay_date,
        "davr": f"{sh.start:%d.%m.%Y} — {sh.end:%d.%m.%Y}",
        "davr_boshi": sh.start,
        "davr_oxiri": sh.end,
        "avans": float(sh.total),
        "avans_soni": len(sh.rows),
    }


def _row_values(i: int, t) -> dict:
    usd = t.currency == "USD" and t.orig_amount is not None
    return {
        "avans_n": i,
        "avans_sana": t.date,
        "avans_summa": float(t.amount),
        "avans_kim": t.member.name,
        "avans_izoh": safe_cell(t.comment or "") or f"Avans (berdi: {t.member.name})",
        "avans_valyuta": "USD" if usd else "UZS",
        "avans_asl": float(t.orig_amount) if usd else float(t.amount),
        "avans_kurs": float(t.rate) if usd and t.rate else None,
    }


def _rest(rows: list) -> SimpleNamespace:
    """Advances that don't fit the template's table: one line, in so'm."""
    return SimpleNamespace(
        date=rows[-1].date,
        amount=sum((t.amount for t in rows), S.ZERO),
        currency="UZS",
        orig_amount=None,
        rate=None,
        member=rows[-1].member,
        comment=f"+{len(rows)} ta avans birga — ro'yxati «Avanslar» varag'ida",
    )


def _set(cell, x) -> None:
    cell.value = x
    if isinstance(x, date) and cell.number_format in ("General", ""):
        cell.number_format = "DD.MM.YYYY"


def _fill(cell, values: dict) -> None:
    v = cell.value
    if not isinstance(v, str) or "{" not in v:
        return
    whole = PH.fullmatch(v.strip())
    if whole and whole.group(1) in values:
        x = values[whole.group(1)]
        _set(cell, "" if x is None else x)
        if isinstance(x, float) and cell.number_format in ("General", ""):
            cell.number_format = "#,##0"
        return

    def rep(m):
        x = values.get(m.group(1), m.group(0))
        if isinstance(x, date):
            return f"{x:%d.%m.%Y}"
        if isinstance(x, float):
            return f"{x:,.0f}".replace(",", " ")
        return "" if x is None else str(x)

    cell.value = PH.sub(rep, v)


def _has_placeholders(ws) -> bool:
    return any(
        isinstance(c.value, str) and PH.search(c.value) for row in ws.iter_rows() for c in row
    )


def _fit(rows: list, room: int) -> list:
    if len(rows) <= room:
        return rows
    return rows[: room - 1] + [_rest(rows[room - 1 :])] if room > 1 else [_rest(rows)]


def _placeholder_list(ws, rows: list) -> None:
    """The row with {avans_…}: repeated downwards — only its placeholder cells, so
    the template's own formulas and borders stay as they are."""
    r0 = next(
        (
            c.row
            for row in ws.iter_rows()
            for c in row
            if isinstance(c.value, str) and any(k in c.value for k in ROW_KEYS)
        ),
        None,
    )
    if r0 is None:
        return
    tpl = {
        c.column: (c.value, c._style)
        for c in ws[r0]
        if isinstance(c.value, str) and any(k in c.value for k in ROW_KEYS)
    }
    if not rows:
        for col in tpl:
            _fill(ws.cell(r0, col), {k: "" for k in ROW_KEYS})
        return
    for i, t in enumerate(rows):
        vals = _row_values(i + 1, t)
        for col, (value, style) in tpl.items():
            c = ws.cell(r0 + i, col)
            if i:
                c._style = copy(style)
                c.value = value
            _fill(c, vals)


# Templates without placeholders (the ones on file: ODDIY, OFIS, MENEJER, ZAMER,
# SBORKA): the labels on the left and the «… AVANSLAR» table are found by their text.
def _text(c) -> str:
    return " ".join(str(c.value).split()).lower() if isinstance(c.value, str) else ""


def _merged_end(ws, c) -> int:
    for r in ws.merged_cells.ranges:
        if c.coordinate in r:
            return r.max_col
    return c.column


def _value_cell(ws, label):
    return ws.cell(label.row, _merged_end(ws, label) + 1)


def _empty(c) -> bool:
    return c.value is None or (isinstance(c.value, str) and not c.value.strip())


HEAD_LABELS = [
    (re.compile(r"^xodim( \(f\.?i\.?sh\.?\))?$"), "ism"),
    (re.compile(r"^lavozim$"), "lavozim"),
    (re.compile(r"^(soha|bo'lim)( / (soha|bo'lim))?$"), "bolim"),
    (re.compile(r"^davr \(oy 1-sanasi\)$"), "oy_boshi"),
    (re.compile(r"^davr$"), "oy"),
    (re.compile(r"^hisob sanasi$"), "sana"),
]
RATE_LABEL = re.compile(r"^1 usd kursi")


def _labels(ws, values: dict, rate: Decimal | None) -> None:
    for row in ws.iter_rows(max_row=min(ws.max_row, 12)):
        for c in row:
            t = _text(c)
            if not t:
                continue
            if rate and RATE_LABEL.match(t):
                v = _value_cell(ws, c)
                if not isinstance(v.value, str):  # a number or empty: today's rate
                    v.value = float(rate)
                continue
            for rx, key in HEAD_LABELS:
                if rx.match(t):
                    v = _value_cell(ws, c)
                    if _empty(v) and values.get(key) not in (None, ""):
                        _set(v, values[key])
                    break


AVANS_TITLE = re.compile(r"avans")
COLS = {
    "sana": "avans_sana",
    "izoh": "avans_izoh",
    "kim oldi": "kim",
    "nomi": "nomi",
    "valyuta": "avans_valyuta",
    "summa": "avans_asl",
    "kurs": "avans_kurs",
    "kurs*": "avans_kurs",
}


def _advance_table(ws):
    """(header row, {column: key}, [data rows]) of the advances table, or None."""
    rows = list(ws.iter_rows())
    for i, row in enumerate(rows):
        if not any(AVANS_TITLE.search(_text(c)) and not _text(c).startswith("jami") for c in row):
            continue
        for head in rows[i + 1 : i + 3]:
            cols = {}
            for c in head:
                t = _text(c)
                if t in COLS:
                    cols[c.column] = COLS[t]
                elif t.startswith("summa") and "avans_asl" not in cols.values():
                    cols[c.column] = "som"  # «Summa (so'm)»: a table without currency
            if "avans_sana" not in cols.values():
                continue
            if "avans_valyuta" not in cols.values():
                cols = {k: ("som" if v == "avans_asl" else v) for k, v in cols.items()}
            data = []
            for r in range(head[0].row + 1, ws.max_row + 1):
                # data rows are numbered in column A (=ROW()-N or 1, 2, 3 …)
                first = ws.cell(r, 1).value
                numbered = isinstance(first, (int, float)) or (
                    isinstance(first, str) and first.startswith("=") and "row()" in first.lower()
                )
                if not numbered:
                    break
                data.append(r)
            if data:
                return head[0].row, cols, data
    return None


def _auto_list(ws, rows: list, name: str) -> bool:
    found = _advance_table(ws)
    if not found:
        return False
    _, cols, data = found
    free = [r for r in data if all(_empty(ws.cell(r, col)) for col in cols)]
    for r, t in zip(free, _fit(rows, len(free))):
        vals = {
            **_row_values(0, t),
            "kim": name,
            "nomi": "Avans",
            "som": float(t.amount),
        }
        for col, key in cols.items():
            x = vals.get(key)
            if x is not None:
                _set(ws.cell(r, col), x)
        ws.row_dimensions[r].hidden = False
    return True


def _list_sheet(wb, sh: Sheet) -> None:
    name = "Avanslar"
    if name in wb.sheetnames:
        del wb[name]
    ws = wb.create_sheet(name)
    bold = Font(bold=True)
    ws["A1"] = f"{sh.worker.name} — avanslar ({sh.start:%d.%m.%Y} — {sh.end:%d.%m.%Y})"
    ws["A1"].font = Font(bold=True, size=12)
    head = PatternFill("solid", fgColor="FFF3E0")
    for c, h in enumerate(["№", "Sana", "Summa (so'm)", "Dollarda", "Kim berdi", "Izoh"], start=1):
        x = ws.cell(3, c, h)
        x.font, x.fill = bold, head
    for i, t in enumerate(sh.rows, start=1):
        r = 3 + i
        ws.cell(r, 1, i)
        ws.cell(r, 2, t.date).number_format = "DD.MM.YYYY"
        ws.cell(r, 3, float(t.amount)).number_format = "#,##0"
        if t.currency == "USD" and t.orig_amount is not None:
            ws.cell(r, 4, float(t.orig_amount)).number_format = "#,##0.##"
        ws.cell(r, 5, t.member.name)
        ws.cell(r, 6, safe_cell(t.comment or "")).alignment = Alignment(wrap_text=True)
    r = 4 + len(sh.rows)
    ws.cell(r, 1, "Jami").font = bold
    ws.cell(r, 3, float(sh.total)).number_format = "#,##0"
    ws.cell(r, 3).font = bold
    for col, w in zip("ABCDEF", (6, 12, 16, 12, 22, 48)):
        ws.column_dimensions[col].width = w


def _rate() -> Decimal | None:
    try:
        r = S.usd_rate()
    except Exception:  # noqa: BLE001 — the sheet goes out with the template's rate
        return None
    return r if r and r > 0 else None


def salary_xlsx(
    worker: Worker, ym: str | None = None, start: date | None = None, end: date | None = None
) -> tuple[bytes, str]:
    """The worker's template with the name and advances of `ym` filled in."""
    sh = salary_info(worker, ym, start, end)
    content, _ = template_file(worker)
    wb = load_workbook(io.BytesIO(content))
    values = _values(sh)
    rate = _rate()
    for ws in wb.worksheets:
        if _has_placeholders(ws):
            _placeholder_list(ws, sh.rows)
            for row in ws.iter_rows():
                for c in row:
                    _fill(c, values)
            continue
        _labels(ws, values, rate)
        name_cell = next(
            (
                _value_cell(ws, c)
                for row in ws.iter_rows(max_row=min(ws.max_row, 12))
                for c in row
                if HEAD_LABELS[0][0].match(_text(c))
            ),
            None,
        )
        name = name_cell.value if name_cell is not None and name_cell.value else worker.name
        _auto_list(ws, sh.rows, str(name))
    _list_sheet(wb, sh)
    wb.calculation.fullCalcOnLoad = True  # totals / OSTATOK recalculated on opening
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), f"Oylik_{sh.ym}_{_slug(worker.name)}.xlsx"


# ------------------------------------------------------------------ reminder (celery)
def due_tomorrow(on: date | None = None) -> list[tuple[Worker, date]]:
    """Active workers whose pay day is tomorrow."""
    t = (on or today()) + timedelta(days=1)
    ym = t.strftime("%Y-%m")
    return [
        (w, t)
        for w in Worker.objects.filter(is_active=True, pay_day__isnull=False)
        if pay_date(w, ym) == t
    ]


def remind(on: date | None = None) -> int:
    """The day before a pay day: «{name} — oylik chiqarish kerak» to the accountant.
    Once per worker and pay date (beat restarts and redelivery send nothing twice)."""
    due = [(w, d) for w, d in due_tomorrow(on) if S.once(f"salary:{w.pk}:{d.isoformat()}")]
    if due:
        E.salary_due(due)
    return len(due)

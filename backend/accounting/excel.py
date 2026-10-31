"""«📥 Excel»: the month report file (what the «Xulosa» and «Tekshiruv» sheets showed)."""

from __future__ import annotations

import io
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import services as S
from .constants import CAT_SALARY, OpKind
from .models import MonthReport
from .text import add_months, month_bounds, month_full, month_name, safe_cell

DARK, MINT, BAND = "0B3D2E", "E8F3EC", "F3F8F5"
HEAD_FONT = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor=DARK)
MONEY = "#,##0;[Red]-#,##0"


def _head(ws, row, values):
    for i, v in enumerate(values, 1):
        c = ws.cell(row=row, column=i, value=v)
        c.font, c.fill = HEAD_FONT, HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center")


def _section(ws, row, text):
    ws.cell(row=row, column=1, value=text).font = Font(bold=True, color=DARK, size=12)


def _widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build_xlsx(start: date, end: date, title: str = "") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Xulosa"
    sm = S.summary(start, end)
    ws["A1"] = f"📊 SerMebel — Xulosa {title}".strip()
    ws["A1"].font = Font(bold=True, size=16, color=DARK)
    ws["A2"] = f"Davr: {start:%d.%m.%Y} — {end:%d.%m.%Y}"
    cards = [
        ("Xarajatlar", -sm["expense"]),
        ("Daromadlar", sm["income"]),
        ("Balans", sm["balance"]),
        ("Xodimlarga berildi", sm["given_to_staff"]),
        ("Rahbarga (shaxsiy)", sm["boss_personal"]),
        ("Uy uchun (shaxsiy)", sm["home"]),
    ]
    for i, (lab, val) in enumerate(cards):
        ws.cell(row=4, column=1 + i * 2, value=lab).font = Font(bold=True, color=DARK)
        c = ws.cell(row=5, column=1 + i * 2, value=float(val))
        c.number_format, c.font = MONEY, Font(bold=True, size=13)
    r = 7
    _section(ws, r, "Xarajatlar turlari bo'yicha")
    _head(ws, r + 1, ["Kategoriya", "Summa", "Operatsiya", "Ulushi"])
    r += 2
    tot = sum(x["sum"] for x in sm["expense_by_category"]) or 1
    for x in sm["expense_by_category"]:
        ws.append([f"{x['emoji']} {x['name']}", float(x["sum"]), x["n"], float(x["sum"] / tot)])
        ws.cell(row=r, column=2).number_format = MONEY
        ws.cell(row=r, column=4).number_format = "0%"
        r += 1
    r += 1
    _section(ws, r, "Daromadlar turlari bo'yicha")
    _head(ws, r + 1, ["Turi", "Summa", "Operatsiya"])
    r += 2
    for x in sm["income_by_category"]:
        ws.append([f"{x['emoji']} {x['name']}", float(x["sum"]), x["n"]])
        ws.cell(row=r, column=2).number_format = MONEY
        r += 1
    r += 1
    _section(ws, r, "Hamyonlar (xodimlar)")
    _head(
        ws, r + 1, ["Xodim", "Davrda berildi", "Davrda sarfladi", "Hamyonda (jami)", "Oylik hisob"]
    )
    r += 2
    bal = S.balances()
    ym = start.strftime("%Y-%m")
    reported = set(MonthReport.objects.filter(month=ym).values_list("member_id", flat=True))
    for m in S.wallets():
        ops = list(S.period_ops(start, end, m))
        given = sum(float(t.amount) for t in ops if t.kind == OpKind.INCOME)
        spent = sum(
            float(t.amount)
            for t in ops
            if t.kind == OpKind.EXPENSE and t.category.code != "otkazma"
        )
        ws.append(
            [
                ("💼 " if m.is_admin else "👤 ") + m.name,
                given,
                spent,
                "— kassa" if m.is_admin else float(bal.get(m.id, 0)),
                "—" if m.is_admin else ("✅ topshirdi" if m.id in reported else "⏳ yo'q"),
            ]
        )
        for col in (2, 3, 4):
            ws.cell(row=r, column=col).number_format = MONEY
        r += 1
    r += 1
    _section(ws, r, "👷 Oylik — kimga")
    _head(ws, r + 1, ["Ism", "Davrda berildi"])
    r += 2
    for name, val in sm["salary_by_worker"]:
        ws.append([name, float(val)])
        ws.cell(row=r, column=2).number_format = MONEY
        r += 1
    r += 1
    _section(ws, r, "Oylar bo'yicha (12 oy)")
    _head(ws, r + 1, ["Oy", "Daromad", "Xarajat", "Balans"])
    r += 2
    for k in range(-11, 1):
        mm = add_months(ym, k)
        s, e = month_bounds(mm)
        x = S.summary(s, e)
        ws.append([month_name(mm), float(x["income"]), float(x["expense"]), float(x["balance"])])
        for col in (2, 3, 4):
            ws.cell(row=r, column=col).number_format = MONEY
        r += 1
    _widths(ws, [34, 18, 16, 18, 18, 16, 16, 16, 16, 16, 16, 16])

    ws2 = wb.create_sheet("Amallar")
    _head(
        ws2,
        1,
        [
            "№",
            "Sana",
            "Vaqt",
            "Xodim",
            "Turi",
            "Summa (so'm)",
            "Valyuta",
            "Asl summa",
            "Kurs",
            "Kategoriya",
            "Buyurtma ID",
            "Izoh",
            "Ishchi / kimga",
            "Ish haqi turi",
            "Manba",
        ],
    )
    for i, t in enumerate(S.period_ops(start, end).order_by("date", "number"), 2):
        ws2.append(
            [
                t.number,
                t.date,
                t.time.strftime("%H:%M") if t.time else "",
                t.member.name,
                t.get_kind_display(),
                float(t.amount),
                t.currency,
                float(t.orig_amount) if t.orig_amount else None,
                float(t.rate) if t.rate else None,
                t.category.name,
                t.order_ref,
                safe_cell(t.comment),
                safe_cell(t.worker),
                t.pay_type if t.category.code == CAT_SALARY else "",
                t.get_source_display(),
            ]
        )
        ws2.cell(row=i, column=2).number_format = "dd.mm.yyyy"
        ws2.cell(row=i, column=6).number_format = MONEY
    ws2.freeze_panes = "A2"
    _widths(ws2, [8, 12, 8, 22, 10, 16, 8, 12, 10, 26, 12, 40, 24, 12, 12])

    ws3 = wb.create_sheet("Tekshiruv")
    _head(ws3, 1, ["Daraja", "Muammo", "Tafsilot", "Yechim"])
    lv = {"h": "🔴", "m": "🟠", "l": "🟡"}
    issues = S.checks()
    for i in issues:
        ws3.append([lv[i["level"]], i["title"], i["detail"], i["fix"]])
    if not issues:
        ws3.append(["", "✅ Kamchilik topilmadi"])
    for row in ws3.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    _widths(ws3, [8, 48, 52, 60])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def month_xlsx(ym: str) -> tuple[bytes, str]:
    s, e = month_bounds(ym)
    return build_xlsx(s, e, month_full(ym)), f"SerMebel_{ym}.xlsx"


# ------------------------------------------------------------------ one person's month
def member_month_xlsx(member, ym: str) -> tuple[bytes, str]:
    """The xlsx a staff member's «📤 Oylik hisob» comes with: totals, by category, by order,
    salaries paid, and every operation of the month with the running balance."""
    s, e = month_bounds(ym)
    sm = S.month_summary(member, ym)
    rep = MonthReport.objects.filter(member=member, month=ym).first()
    ops = list(S.period_ops(s, e, member).order_by("date", "time", "number"))
    inc = [t for t in ops if t.kind == OpKind.INCOME]
    out = [t for t in ops if t.kind == OpKind.EXPENSE]

    wb = Workbook()
    ws = wb.active
    ws.title = "Hisobot"
    ws["A1"] = f"📤 {member.name} — {month_full(ym)} hisobi"
    ws["A1"].font = Font(bold=True, size=16, color=DARK)
    ws["A2"] = f"Davr: {s:%d.%m.%Y} — {e:%d.%m.%Y}"
    if rep:
        ws["A3"] = f"Topshirildi: {rep.submitted_at:%d.%m.%Y %H:%M}" + (
            f" · Izoh: {safe_cell(rep.note)}" if rep.note else ""
        )
    cards = [
        ("Oy boshida qoldiq", sm["opening"]),
        ("Olindi (kirim)", sm["income"]),
        ("Sarflandi (chiqim)", sm["expense"]),
        ("Oy oxirida qoldiq", sm["closing"]),
        ("Amallar soni", sm["count"]),
    ]
    for i, (lab, val) in enumerate(cards):
        ws.cell(row=5, column=1 + i, value=lab).font = Font(bold=True, color=DARK)
        c = ws.cell(row=6, column=1 + i, value=float(val))
        c.font = Font(bold=True, size=13)
        c.number_format = "0" if lab == "Amallar soni" else MONEY
        c.fill = PatternFill("solid", fgColor=MINT)

    def grouped(rows, key):
        g: dict[str, list] = {}
        for t in rows:
            k = key(t)
            x = g.setdefault(k, [0.0, 0])
            x[0] += float(t.amount)
            x[1] += 1
        return sorted(g.items(), key=lambda kv: -kv[1][0])

    def table(r, title, head, rows, share=False):
        _section(ws, r, title)
        _head(ws, r + 1, head)
        r += 2
        tot = sum(v[0] for _, v in rows) or 1
        if not rows:
            ws.cell(row=r, column=1, value="—")
            return r + 2
        for name, (val, n) in rows:
            vals = [name, val, n] + ([val / tot] if share else [])
            for col, v in enumerate(vals, 1):
                ws.cell(row=r, column=col, value=v)
            ws.cell(row=r, column=2).number_format = MONEY
            if share:
                ws.cell(row=r, column=4).number_format = "0%"
            r += 1
        ws.cell(row=r, column=1, value="Jami").font = Font(bold=True)
        c = ws.cell(row=r, column=2, value=sum(v[0] for _, v in rows))
        c.number_format, c.font = MONEY, Font(bold=True)
        ws.cell(row=r, column=3, value=sum(v[1] for _, v in rows)).font = Font(bold=True)
        return r + 2

    def cat(t):
        return f"{t.category.emoji} {t.category.name}".strip()

    r = table(
        8,
        "🔴 Xarajatlar turlari bo'yicha",
        ["Kategoriya", "Summa", "Amal", "Ulushi"],
        grouped(out, cat),
        share=True,
    )
    r = table(r, "🟢 Kirimlar (kimdan / nima)", ["Turi", "Summa", "Amal"], grouped(inc, cat))
    r = table(
        r,
        "🧾 Buyurtmalar bo'yicha xarajat",
        ["Buyurtma", "Summa", "Amal"],
        grouped(out, lambda t: t.order_ref or "— buyurtmasiz"),
    )
    salary = [t for t in out if t.category.code == CAT_SALARY]
    if salary:
        table(
            r,
            "👷 Ishchilarga berilgan oylik",
            ["Ishchi", "Summa", "Amal"],
            grouped(salary, lambda t: t.worker or "— (ism yozilmagan)"),
        )
    _widths(ws, [34, 18, 18, 18, 16])

    ws2 = wb.create_sheet("Amallar")
    _head(
        ws2,
        1,
        [
            "№",
            "Sana",
            "Vaqt",
            "Kirim",
            "Chiqim",
            "Qoldiq",
            "Kategoriya",
            "Buyurtma",
            "Izoh",
            "Ishchi / kimga",
            "Valyuta",
            "Asl summa",
            "Kurs",
            "Chek",
            "Manba",
        ],
    )
    ws2.append(
        ["", s - timedelta(days=1), "", None, None, float(sm["opening"]), "Oy boshida qoldiq"]
    )
    ws2.cell(row=2, column=2).number_format = "dd.mm.yyyy"
    ws2.cell(row=2, column=6).number_format = MONEY
    for c in ws2[2]:
        c.font = Font(italic=True)
    run = float(sm["opening"])
    for i, t in enumerate(ops, 3):
        a = float(t.amount)
        is_in = t.kind == OpKind.INCOME
        run += a if is_in else -a
        ws2.append(
            [
                t.number,
                t.date,
                t.time.strftime("%H:%M") if t.time else "",
                a if is_in else None,
                None if is_in else a,
                run,
                cat(t),
                t.order_ref,
                safe_cell(t.comment),
                safe_cell(t.worker),
                t.currency,
                float(t.orig_amount) if t.orig_amount else None,
                float(t.rate) if t.rate else None,
                "✅" if t.receipt_file_id else "",
                t.get_source_display(),
            ]
        )
        ws2.cell(row=i, column=2).number_format = "dd.mm.yyyy"
        for col in (4, 5, 6):
            ws2.cell(row=i, column=col).number_format = MONEY
        if i % 2:
            for c in ws2[i]:
                c.fill = PatternFill("solid", fgColor=BAND)
    last = len(ops) + 3
    ws2.cell(row=last, column=3, value="Jami").font = Font(bold=True)
    for col, val in ((4, sm["income"]), (5, sm["expense"]), (6, sm["closing"])):
        c = ws2.cell(row=last, column=col, value=float(val))
        c.number_format, c.font = MONEY, Font(bold=True)
    ws2.freeze_panes = "A2"
    _widths(ws2, [8, 12, 8, 16, 16, 16, 26, 12, 40, 22, 8, 12, 10, 6, 10])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), f"Hisobot_{ym}_{member.pk}.xlsx"

"""Ombor Excel: current stock + the movements of a period."""

from __future__ import annotations

import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from accounting.finance.text import safe_cell as T

from .models import Movement, Product


def _head(ws, values, widths):
    ws.append(values)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0B3D2E")
    ws.freeze_panes = "A2"
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w


def build_xlsx(start: date | None, end: date | None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Qoldiq"
    _head(
        ws,
        ["Mahsulot", "Kod", "Turi", "Birlik", "Qoldiq", "O'rtacha tannarx", "Qiymati", "Minimal"],
        [36, 14, 18, 10, 12, 18, 18, 10],
    )
    for p in Product.objects.filter(is_active=True).order_by("name"):
        ws.append(
            [
                T(p.name),
                T(p.code),
                T(p.category),
                p.get_unit_display(),
                float(p.quantity),
                float(p.avg_cost),
                float(p.quantity * p.avg_cost),
                float(p.min_quantity),
            ]
        )
    mv = Movement.objects.filter(is_cancelled=False).select_related(
        "product", "created_by", "order"
    )
    if start:
        mv = mv.filter(date__gte=start)
    if end:
        mv = mv.filter(date__lte=end)
    ws2 = wb.create_sheet("Harakatlar")
    _head(
        ws2,
        [
            "№",
            "Sana",
            "Turi",
            "Mahsulot",
            "Miqdor",
            "Valyuta",
            "Narx ($)",
            "Kurs",
            "Narx (so'm)",
            "Summa (so'm)",
            "Kimdan / kimga",
            "Buyurtma",
            "Izoh",
            "Kiritdi",
        ],
        [8, 12, 8, 32, 10, 8, 12, 10, 14, 16, 24, 10, 32, 18],
    )
    for m in mv.order_by("date", "number"):
        ws2.append(
            [
                m.number,
                m.date,
                m.get_kind_display(),
                T(m.product.name),
                float(m.quantity),
                "$" if m.currency == "USD" else "so'm",
                float(m.orig_price) if m.orig_price is not None else "",
                float(m.rate) if m.rate else "",
                float(m.unit_price),
                float(m.total),
                T(m.counterparty),
                m.order.order_number if m.order else "",
                T(m.comment),
                T(m.created_by.name or str(m.created_by.phone)) if m.created_by else "",
            ]
        )
        ws2.cell(row=ws2.max_row, column=2).number_format = "dd.mm.yyyy"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

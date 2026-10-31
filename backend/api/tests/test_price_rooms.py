"""Price calculation by rooms: rooms → furniture, the «Excel» sheet, the offer PDF."""

import io

import pytest
from djmoney.money import Money
from openpyxl import load_workbook
from PIL import Image

from metering.price.models import Calculate, Inventory, InventoryType, ObjectType, Price, RoomType


@pytest.fixture
def catalog(db):
    kitchen = ObjectType.objects.create(name="Кухня", count_name="pog")
    wardrobe = ObjectType.objects.create(name="Шкаф", count_name="kv")
    facade = InventoryType.objects.create(
        name="Фасад материал", obj=kitchen, type=InventoryType.TYPE_KV
    )
    mech = InventoryType.objects.create(name="Механизм", obj=kitchen, type=InventoryType.TYPE_COUNT)
    return {
        "kitchen": kitchen,
        "wardrobe": wardrobe,
        "facade": facade,
        "mech": mech,
        "mdf": Inventory.objects.create(name="МДФ краш", type=facade, price=Money(300, "USD")),
        "blum": Inventory.objects.create(name="BLUM", type=mech, price=Money(10, "USD")),
        "oshxona": _room("Oshxona", "🍳"),
        "yotoq": _room("Yotoqxona", "🛏"),
    }


def _room(name, icon):
    """Migration 0017 already seeds the room types: reuse it, with an empty list
    (= every piece of furniture), as the tests expect."""
    room, _ = RoomType.objects.update_or_create(name=name, defaults={"icon": icon})
    room.object_types.clear()
    return room


@pytest.fixture
def price(db):
    return Price.objects.create(metering=None, desc="Zilola opa")


def url(price, tail=""):
    return f"/api/v1/prices/{price.pk}/{tail}"


def add_kitchen(api_client, price, room_id, catalog, count=4.0):
    return api_client.post(
        url(price, "calculations/"),
        {
            "name": "Кухня",
            "count": count,
            "object_type": catalog["kitchen"].pk,
            "room": room_id,
            "note": "Подсветка",
            "items": [
                {"inventory_type": catalog["facade"].pk, "inventory": catalog["mdf"].pk},
                {"inventory_type": catalog["mech"].pk, "inventory": catalog["blum"].pk, "count": 5},
            ],
        },
        format="json",
    )


def test_rooms_are_numbered_and_hold_furniture(admin_api, price, catalog):
    r = admin_api.post(url(price, "rooms/"), {"room_type": catalog["yotoq"].pk}, format="json")
    assert r.status_code == 200, r.data
    r = admin_api.post(url(price, "rooms/"), {"room_type": catalog["yotoq"].pk}, format="json")
    r = admin_api.post(url(price, "rooms/"), {"room_type": catalog["oshxona"].pk}, format="json")
    names = [room["name"] for room in r.data["rooms"]]
    assert names == ["Yotoqxona", "Yotoqxona 2", "Oshxona"]

    kitchen_room = r.data["rooms"][2]["id"]
    r = add_kitchen(admin_api, price, kitchen_room, catalog)
    assert r.status_code == 200, r.data
    calc = r.data["calculations"][0]
    assert calc["room"] == kitchen_room and calc["note"] == "Подсветка"
    # 4 п.м. × 300 + 5 × 10
    assert r.data["price"]["amount"] == 1250

    sheet = r.data["sheet"]
    room = next(x for x in sheet["rooms"] if x["id"] == kitchen_room)
    item = room["items"][0]
    assert item["count_text"] == "4 п.м."
    assert {"label": "Фасад материал", "value": "МДФ краш"} in item["params"]
    assert {"label": "Механизм", "value": "BLUM × 5"} in item["params"]
    assert room["total"] == 1250 and sheet["total"] == 1250
    assert sheet["client"] == "Zilola opa"


def test_room_type_limits_furniture(admin_api, price, catalog):
    catalog["yotoq"].object_types.set([catalog["wardrobe"]])
    r = admin_api.post(url(price, "rooms/"), {"room_type": catalog["yotoq"].pk}, format="json")
    r = add_kitchen(admin_api, price, r.data["rooms"][0]["id"], catalog)
    assert r.status_code == 400
    assert "mos emas" in r.data["detail"]


def test_room_of_another_price_is_refused(admin_api, price, catalog):
    other = Price.objects.create(metering=None)
    r = admin_api.post(url(other, "rooms/"), {"room_type": catalog["oshxona"].pk}, format="json")
    r = add_kitchen(admin_api, price, r.data["rooms"][0]["id"], catalog)
    assert r.status_code == 404


def test_deleting_a_room_removes_its_furniture_and_total(admin_api, price, catalog):
    r = admin_api.post(url(price, "rooms/"), {"room_type": catalog["oshxona"].pk}, format="json")
    room_id = r.data["rooms"][0]["id"]
    add_kitchen(admin_api, price, room_id, catalog)
    r = admin_api.delete(url(price, f"rooms/{room_id}/"))
    assert r.status_code == 200
    assert r.data["rooms"] == [] and r.data["calculations"] == []
    assert r.data["price"]["amount"] == 0
    assert not Calculate.objects.filter(price=price).exists()


def test_rename_and_reorder_rooms(admin_api, price, catalog):
    admin_api.post(url(price, "rooms/"), {"room_type": catalog["oshxona"].pk}, format="json")
    r = admin_api.post(url(price, "rooms/"), {"room_type": catalog["yotoq"].pk}, format="json")
    a, b = [x["id"] for x in r.data["rooms"]]
    r = admin_api.patch(url(price, f"rooms/{b}/"), {"name": "Master-spalnya"}, format="json")
    assert r.data["rooms"][1]["name"] == "Master-spalnya"
    r = admin_api.post(url(price, "rooms/reorder/"), {"ids": [b, a]}, format="json")
    assert [x["id"] for x in r.data["rooms"]] == [b, a]
    assert (
        admin_api.post(url(price, "rooms/reorder/"), {"ids": [b]}, format="json").status_code == 400
    )


def test_old_calculations_without_room_show_as_other(admin_api, price, catalog):
    Calculate.objects.create(
        name="Eski", price=price, obj=catalog["kitchen"], count=1, amount=Money(100, "USD")
    )
    sheet = admin_api.get(url(price)).data["sheet"]
    assert sheet["rooms"][-1]["name"] == "Boshqa"
    assert sheet["total"] == 100


def test_viewer_cannot_edit_rooms(api, price, catalog):
    client = api("price.view_price")
    r = client.post(url(price, "rooms/"), {"room_type": catalog["oshxona"].pk}, format="json")
    assert r.status_code == 403
    assert client.get("/api/v1/room-types/").status_code == 200


def test_excel_and_pdf(admin_api, price, catalog, settings, tmp_path):
    settings.MEDIA_ROOT = tmp_path
    r = admin_api.post(url(price, "rooms/"), {"room_type": catalog["oshxona"].pk}, format="json")
    room_id = r.data["rooms"][0]["id"]
    add_kitchen(admin_api, price, room_id, catalog)

    buf = io.BytesIO()
    Image.new("RGB", (400, 300), "teal").save(buf, "PNG")
    buf.seek(0)
    buf.name = "render.png"
    r = admin_api.post(url(price, f"rooms/{room_id}/images/"), {"file": buf}, format="multipart")
    assert r.status_code == 200, r.data
    assert len(r.data["rooms"][0]["images"]) == 1

    xlsx = admin_api.get(url(price, "excel/"))
    assert xlsx.status_code == 200
    flat = [
        str(c)
        for row in load_workbook(io.BytesIO(xlsx.content)).active.iter_rows(values_only=True)
        for c in row
        if c is not None
    ]
    assert "🍳 Oshxona" in flat and "МДФ краш" in flat and "Жами" in flat

    pdf = admin_api.get(url(price, "pdf/"))
    assert pdf.status_code == 200
    assert pdf["Content-Type"] == "application/pdf"
    assert pdf.content.startswith(b"%PDF")

    image_id = r.data["rooms"][0]["images"][0]["id"]
    r = admin_api.delete(url(price, f"rooms/{room_id}/images/{image_id}/"))
    assert r.data["rooms"][0]["images"] == []


def test_non_image_upload_is_refused(admin_api, price, catalog, settings, tmp_path):
    settings.MEDIA_ROOT = tmp_path
    r = admin_api.post(url(price, "rooms/"), {"room_type": catalog["oshxona"].pk}, format="json")
    room_id = r.data["rooms"][0]["id"]
    buf = io.BytesIO(b"%PDF-1.4 fake")
    buf.name = "doc.pdf"
    r = admin_api.post(url(price, f"rooms/{room_id}/images/"), {"file": buf}, format="multipart")
    assert r.status_code == 400


def test_offer_title_and_variant_are_editable(admin_api, price):
    r = admin_api.patch(
        url(price), {"client_title": "Зилола опа", "variant": "Премиум"}, format="json"
    )
    assert r.status_code == 200
    assert r.data["sheet"]["client"] == "Зилола опа"
    assert r.data["sheet"]["variant"] == "Премиум"


def test_calculation_delete_works_but_price_delete_does_not(admin_api, price, catalog):
    r = admin_api.post(url(price, "rooms/"), {"room_type": catalog["oshxona"].pk}, format="json")
    r = add_kitchen(admin_api, price, r.data["rooms"][0]["id"], catalog)
    calc_id = r.data["calculations"][0]["id"]
    r = admin_api.delete(url(price, f"calculations/{calc_id}/"))
    assert r.status_code == 200 and r.data["calculations"] == []
    assert admin_api.delete(url(price)).status_code == 405
    assert Price.objects.filter(pk=price.pk).exists()


# --- colour, accessories, manual price, versions, the detailing lock ---------------


def _png(name="acc.png"):
    buf = io.BytesIO()
    Image.new("RGB", (300, 300), "orange").save(buf, "PNG")
    buf.seek(0)
    buf.name = name
    return buf


def test_accessories_colour_and_manual_price(admin_api, price, catalog, settings, tmp_path):
    settings.MEDIA_ROOT = tmp_path
    r = admin_api.post(url(price, "rooms/"), {"room_type": catalog["oshxona"].pk}, format="json")
    room_id = r.data["rooms"][0]["id"]
    img = admin_api.post(url(price, "accessory-images/"), {"file": _png()}, format="multipart")
    assert img.status_code == 200, img.data
    body = {
        "name": "Кухня",
        "count": 4,
        "object_type": catalog["kitchen"].pk,
        "room": room_id,
        "color": "Оқ матовий",
        "items": [{"inventory_type": catalog["facade"].pk, "inventory": catalog["mdf"].pk}],
        "accessories": [
            {"inventory": catalog["blum"].pk, "count": 3},  # catalog price 10
            {"name": "Ручка золото", "count": 8, "price": "5", "image": img.data["image"]},
        ],
    }
    r = admin_api.post(url(price, "calculations/"), body, format="json")
    assert r.status_code == 200, r.data
    calc = r.data["calculations"][0]
    # 4 × 300 + 3 × 10 + 8 × 5
    assert float(calc["auto_amount"]) == 1270 and r.data["price"]["amount"] == 1270
    assert [a["name"] for a in calc["accessories"]] == ["BLUM", "Ручка золото"]
    assert calc["accessories"][1]["thumbnail"]

    lines = r.data["sheet"]["rooms"][0]["items"][0]["lines"]
    assert {"label": "Ранги", "value": "Оқ матовий"} in lines
    assert {"label": "Аксессуар", "value": "Ручка золото × 8"} in lines
    # from the catalog: the line says what it is (its catalog type), not just «Аксессуар»
    assert {"label": "Механизм", "value": "BLUM × 3"} in lines
    assert calc["accessories"][0]["type_name"] == "Механизм"

    # final price by hand; the accessory rows are updated in place, one removed
    body["manual_amount"] = "1200"
    body["accessories"] = [dict(calc["accessories"][1], price=5)]
    r = admin_api.patch(url(price, f"calculations/{calc['id']}/"), body, format="json")
    assert r.status_code == 200, r.data
    calc2 = r.data["calculations"][0]
    assert r.data["price"]["amount"] == 1200
    assert float(calc2["auto_amount"]) == 1240
    assert [a["id"] for a in calc2["accessories"]] == [calc["accessories"][1]["id"]]

    pdf = admin_api.get(url(price, "pdf/"))
    assert pdf.content.startswith(b"%PDF")


def test_no_colour_no_accessories_no_lines(admin_api, price, catalog):
    r = admin_api.post(url(price, "rooms/"), {"room_type": catalog["oshxona"].pk}, format="json")
    r = add_kitchen(admin_api, price, r.data["rooms"][0]["id"], catalog)
    labels = [ln["label"] for ln in r.data["sheet"]["rooms"][0]["items"][0]["lines"]]
    assert "Ранги" not in labels and "Аксессуар" not in labels


def test_versions_are_kept_and_deduplicated(admin_api, price, catalog):
    r = admin_api.post(url(price, "rooms/"), {"room_type": catalog["oshxona"].pk}, format="json")
    room_id = r.data["rooms"][0]["id"]
    add_kitchen(admin_api, price, room_id, catalog)
    r = admin_api.post(url(price, "versions/"), {}, format="json")
    assert [(v["number"], v["label"], v["total"]) for v in r.data["versions"]] == [
        (1, "Taxminiy narx", 1250)
    ]
    r = admin_api.post(url(price, "versions/"), {}, format="json")  # nothing changed
    assert len(r.data["versions"]) == 1

    add_kitchen(admin_api, price, room_id, catalog, count=5)
    r = admin_api.post(url(price, "versions/"), {}, format="json")
    assert [(v["number"], v["total"]) for v in r.data["versions"]] == [(2, 2800), (1, 1250)]

    v1 = r.data["versions"][1]["id"]
    pdf = admin_api.get(url(price, f"versions/{v1}/pdf/"))
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")


def test_price_is_locked_once_the_order_is_in_detailing(
    admin_api, db_client, catalog, order_factory
):
    from datetime import datetime

    from metering.design.models import Design
    from metering.models import Metering
    from order.constants import OrderStatus

    metering = Metering.objects.create(client=db_client, date_time=datetime(2026, 10, 1, 10))
    Design.objects.get_or_create(metering=metering)
    price = Price.objects.create(metering=metering)
    r = admin_api.post(url(price, "rooms/"), {"room_type": catalog["oshxona"].pk}, format="json")
    room_id = r.data["rooms"][0]["id"]
    order = order_factory(metering=metering)
    assert admin_api.get(url(price)).data["locked"] is False  # «Yaratildi»: still editable

    order.status = OrderStatus.DETAILING
    order.save(update_fields=["status"])
    data = admin_api.get(url(price)).data
    assert data["locked"] is True and data["actions"]["edit"] is False
    for resp in (
        add_kitchen(admin_api, price, room_id, catalog),
        admin_api.post(url(price, "rooms/"), {"room_type": catalog["oshxona"].pk}, format="json"),
        admin_api.delete(url(price, f"rooms/{room_id}/")),
    ):
        assert resp.status_code == 400 and "detalirovka" in resp.data["detail"]
    # the offer can still be saved / printed
    assert admin_api.post(url(price, "versions/"), {}, format="json").status_code == 200


def test_room_picture_can_be_attached_to_furniture(admin_api, price, catalog, settings, tmp_path):
    settings.MEDIA_ROOT = tmp_path
    r = admin_api.post(url(price, "rooms/"), {"room_type": catalog["oshxona"].pk}, format="json")
    room_id = r.data["rooms"][0]["id"]
    r = add_kitchen(admin_api, price, room_id, catalog)
    calc_id = r.data["calculations"][0]["id"]
    # uploaded straight onto the furniture
    r = admin_api.post(
        url(price, f"rooms/{room_id}/images/"),
        {"file": _png("k.png"), "calculate": calc_id},
        format="multipart",
    )
    assert r.status_code == 200, r.data
    img = r.data["rooms"][0]["images"][0]
    assert img["calculate"] == calc_id
    sheet_room = r.data["sheet"]["rooms"][0]
    assert sheet_room["images"] == 0 and sheet_room["items"][0]["images"] == 1
    # detached back to the room, then attached again
    r = admin_api.patch(
        url(price, f"rooms/{room_id}/images/{img['id']}/"), {"calculate": None}, format="json"
    )
    assert r.data["rooms"][0]["images"][0]["calculate"] is None
    r = admin_api.patch(
        url(price, f"rooms/{room_id}/images/{img['id']}/"), {"calculate": calc_id}, format="json"
    )
    assert r.data["rooms"][0]["images"][0]["calculate"] == calc_id
    assert admin_api.get(url(price, "pdf/")).content.startswith(b"%PDF")


def test_unit_and_sizes_give_the_count(admin_api, price, catalog):
    r = admin_api.post(url(price, "rooms/"), {"room_type": catalog["oshxona"].pk}, format="json")
    room_id = r.data["rooms"][0]["id"]
    body = {
        "name": "Шкаф",
        "count": 0,
        "object_type": catalog["kitchen"].pk,  # pog by default
        "room": room_id,
        "unit": "kv",
        "height": 2400,
        "width": 1500,
        "depth": 600,
        "items": [{"inventory_type": catalog["facade"].pk, "inventory": catalog["mdf"].pk}],
    }
    r = admin_api.post(url(price, "calculations/"), body, format="json")
    assert r.status_code == 200, r.data
    calc = r.data["calculations"][0]
    assert calc["count"] == 3.6 and calc["unit"] == "kv"  # 2,4 × 1,5
    assert r.data["price"]["amount"] == 1080  # 3,6 × 300
    item = r.data["sheet"]["rooms"][0]["items"][0]
    assert item["count_text"] == "3,6 м²" and item["size_text"] == "2400 × 1500 × 600 мм"

    body.update(unit="pog", height=None)  # running metres: the width only
    r = admin_api.patch(url(price, f"calculations/{calc['id']}/"), body, format="json")
    assert r.data["calculations"][0]["count"] == 1.5
    assert r.data["sheet"]["rooms"][0]["items"][0]["count_text"] == "1,5 п.м."

    body.update(unit="", width=None, depth=None, count=2)  # no sizes: the typed figure
    r = admin_api.patch(url(price, f"calculations/{calc['id']}/"), body, format="json")
    assert r.data["calculations"][0]["count"] == 2
    assert admin_api.get(url(price, "pdf/")).content.startswith(b"%PDF")


def test_client_pdf_shows_size_and_only_decor_keeps_m2():
    from metering.price.services import client_measure_text

    base = {"count_text": "3,6 м²", "size_text": "2400 × 1500 × 600 мм", "unit": "kv"}
    assert client_measure_text({**base, "show_area": False}) == "2400 × 1500 × 600 мм"
    assert client_measure_text({**base, "show_area": True}) == "3,6 м² · 2400 × 1500 × 600 мм"
    # no size typed: m² is better than nothing
    assert client_measure_text({**base, "size_text": "", "show_area": False}) == "3,6 м²"
    pog = {"count_text": "4,19 п.м.", "size_text": "", "unit": "pog", "show_area": False}
    assert client_measure_text(pog) == "4,19 п.м."


def test_accessory_from_another_furniture_is_refused(admin_api, price, catalog):
    sofa = ObjectType.objects.create(name="Диван", count_name="pog")
    soft = InventoryType.objects.create(name="Материал", obj=sofa, type=InventoryType.TYPE_KV)
    velvet = Inventory.objects.create(name="Велюр", type=soft, price=Money(50, "USD"))
    r = admin_api.post(url(price, "rooms/"), {"room_type": catalog["oshxona"].pk}, format="json")
    body = {
        "name": "Кухня",
        "count": 4,
        "object_type": catalog["kitchen"].pk,
        "room": r.data["rooms"][0]["id"],
        "accessories": [{"inventory": velvet.pk, "count": 1}],
    }
    r = admin_api.post(url(price, "calculations/"), body, format="json")
    assert r.status_code == 400
    assert "alohida mebel" in r.data["detail"]

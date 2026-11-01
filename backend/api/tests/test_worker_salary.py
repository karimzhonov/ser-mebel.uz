"""«Oylik chiqarish»: a worker's own Excel template, the sheet with their advances
filled in, and the accountant's reminder the day before the pay day."""

import io
from datetime import date

import pytest
from openpyxl import Workbook, load_workbook

from accounting import salary as SAL
from accounting import services as S
from accounting.models import Worker

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Oylik: {ism} — {oy}"
    ws["B2"] = "{avans}"
    ws["B3"] = "{{sana}}"
    ws["A5"], ws["B5"], ws["C5"] = "{avans_n}", "{avans_sana}", "{avans_summa}"
    ws["A9"] = "Qo'lda: bonus"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def upload(client, w, content, name="akmal.xlsx"):
    f = io.BytesIO(content)
    f.name = name
    return client.post(f"/api/v1/finance/workers/{w.pk}/template/", {"file": f}, format="multipart")


@pytest.fixture
def worker(members):
    w = Worker.objects.create(group="Sex", name="Akmal", position="Usta", pay_day=5)
    boss = members[0]
    for d, amount in [(date(2026, 10, 5), 1), (date(2026, 10, 6), 300), (date(2026, 10, 20), 200)]:
        S.create_operation(
            member=boss,
            kind="expense",
            amount=amount * 1000,
            category_code="oylik",
            worker="akmal ",
            pay_type="Avans",
            on_date=d,
            actor_admin=True,
        )
    S.create_operation(
        member=boss,
        kind="expense",
        amount=9_000_000,
        category_code="oylik",
        worker="Akmal",
        pay_type="Oylik",
        on_date=date(2026, 10, 25),
        actor_admin=True,
    )
    return w


def test_pay_dates_and_period():
    w = Worker(name="x", pay_day=31)
    assert SAL.pay_date(w, "2026-02") == date(2026, 2, 28)
    assert SAL.period(w, "2026-03") == (date(2026, 3, 1), date(2026, 3, 31))
    w.pay_day = 5
    assert SAL.period(w, "2026-11") == (date(2026, 10, 6), date(2026, 11, 5))
    assert SAL.next_month(w, date(2026, 11, 5)) == "2026-11"
    assert SAL.next_month(w, date(2026, 11, 6)) == "2026-12"
    w.pay_day = None
    assert SAL.period(w, "2026-11") == (date(2026, 11, 1), date(2026, 11, 30))
    # no pay day: the month just ended, the current one from the 25th
    assert SAL.next_month(w, date(2026, 11, 1)) == "2026-10"
    assert SAL.next_month(w, date(2026, 11, 25)) == "2026-11"


def test_upload_and_fill_template(admin_api, worker):
    assert upload(admin_api, worker, b"not excel", "x.xlsx").status_code == 400
    assert upload(admin_api, worker, template(), "x.xls").status_code == 400
    r = upload(admin_api, worker, template())
    assert r.status_code == 200 and r.data["has_template"]
    assert r.data["salary_template_name"] == "akmal.xlsx"

    r = admin_api.get(f"/api/v1/finance/workers/{worker.pk}/salary/?month=2026-11")
    assert r.status_code == 200
    assert (r.data["date_from"], r.data["date_to"]) == (date(2026, 10, 6), date(2026, 11, 5))
    # only advances, only this period (the 5 Oct one belongs to October's salary)
    assert r.data["advance_total"] == 500_000 and len(r.data["advances"]) == 2

    r = admin_api.get(f"/api/v1/finance/workers/{worker.pk}/salary/excel/?month=2026-11")
    assert r.status_code == 200 and r["Content-Type"] == XLSX
    assert "filename*=UTF-8''Oylik_2026-11_Akmal.xlsx" in r["Content-Disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.worksheets[0]
    assert ws["A1"].value.startswith("Oylik: Akmal — ") and "{" not in ws["A1"].value
    assert ws["B2"].value == 500_000
    assert ws["B3"].value.date() == date(2026, 11, 5)
    assert [ws.cell(r, 3).value for r in (5, 6)] == [300_000, 200_000]
    assert ws["A6"].value == 2 and ws["A9"].value == "Qo'lda: bonus"
    assert wb["Avanslar"]["C6"].value == 500_000


def test_starter_template_and_remove(admin_api, worker):
    r = admin_api.get(f"/api/v1/finance/workers/{worker.pk}/template/")
    assert r.status_code == 200 and r["Content-Type"] == XLSX
    assert "{ism}" in [c.value for c in load_workbook(io.BytesIO(r.content)).active["B"]]
    # with no template the sheet still comes out of the starter one
    r = admin_api.get(f"/api/v1/finance/workers/{worker.pk}/salary/excel/?month=2026-11")
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws["B3"].value == "Akmal" and ws["C13"].value == 500_000
    upload(admin_api, worker, template())
    r = admin_api.delete(f"/api/v1/finance/workers/{worker.pk}/template/")
    assert r.status_code == 200 and not r.data["has_template"]


def test_list_does_not_ship_the_file(admin_api, worker):
    upload(admin_api, worker, template())
    r = admin_api.get("/api/v1/finance/workers/")
    assert r.status_code == 200 and "salary_template" not in r.data[0]
    r = admin_api.patch(f"/api/v1/finance/workers/{worker.pk}/", {"pay_day": 40}, format="json")
    assert r.status_code == 400


def test_permissions(api, worker):
    viewer = api("accounting.view_worker")
    assert viewer.get(f"/api/v1/finance/workers/{worker.pk}/salary/").status_code == 403
    pay = api("accounting.view_worker", "accounting.view_all_operations")
    assert pay.get(f"/api/v1/finance/workers/{worker.pk}/salary/excel/").status_code == 200
    assert upload(pay, worker, template()).status_code == 403
    editor = api("accounting.view_worker", "accounting.change_worker")
    assert upload(editor, worker, template()).status_code == 200


def test_worker_card_shows_the_settings(admin_api, worker):
    r = admin_api.get("/api/v1/finance/worker-card/?worker=Akmal&date_from=2026-10-01")
    card = r.data["card"]
    assert card["worker_id"] == worker.pk and card["pay_day"] == 5


def test_reminder_day_before_once(worker, members, bot_events):
    members[1].is_accountant = True
    members[1].save()
    Worker.objects.create(group="Sex", name="Bek", pay_day=6)
    Worker.objects.create(group="Sex", name="Old", pay_day=5, is_active=False)
    assert SAL.remind(date(2026, 11, 3)) == 0
    assert SAL.remind(date(2026, 11, 4)) == 1
    assert SAL.remind(date(2026, 11, 4)) == 0  # beat redelivery: once
    d = bot_events.one("finance.salary_due")
    assert [w["name"] for w in d["workers"]] == ["Akmal"]
    assert d["workers"][0]["pay_date"] == "2026-11-05"
    assert d["to"] == [members[1].telegram_id] and not d["fallback"]


def test_reminder_falls_back_to_admins_and_short_months(members, bot_events):
    Worker.objects.create(group="Sex", name="Oxirgi", pay_day=31)
    assert SAL.remind(date(2026, 2, 27)) == 1
    d = bot_events.one("finance.salary_due")
    assert d["fallback"] and d["to"] == [members[0].telegram_id]
    assert d["workers"][0]["pay_date"] == "2026-02-28"


def office_template() -> bytes:
    """The layout of the office templates (shablonlar/SHABLON_*.xlsx), cut down."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Oylik"
    ws["A2"], ws["C2"] = "Soha / bo'lim", "Ofis"
    ws["A3"], ws["E3"] = "Xodim (F.I.Sh.)", "Davr (oy 1-sanasi)"
    ws["A4"], ws["E4"] = "Lavozim", "Hisob sanasi"
    for r in (2, 3, 4):
        ws.merge_cells(f"A{r}:B{r}")
        ws.merge_cells(f"C{r}:D{r}")
        ws.merge_cells(f"E{r}:F{r}")
    ws["A22"] = "4. OLINGAN AVANSLAR"
    for c, h in zip("ABCEFGH", ["№", "Sana", "Izoh", "Valyuta", "Summa", "Kurs*", "So'm"]):
        ws[f"{c}23"] = h
    for r in (24, 25):
        ws[f"A{r}"] = "=ROW()-23"
        ws[f"H{r}"] = f'=IF(F{r}="",0,IF(E{r}="USD",F{r}*G{r},F{r}))'
    ws["A26"], ws["H26"] = "JAMI AVANSLAR", "=SUM(H24:H25)"
    st = wb.create_sheet("Sozlamalar")
    st["A2"], st["B2"] = "1 USD kursi (so'm)", 11800
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_office_template_is_filled_by_its_labels(admin_api, worker, members, monkeypatch):
    monkeypatch.setattr(S, "usd_rate", lambda: S.D(12500))
    S.create_operation(
        member=members[0],
        kind="expense",
        amount=1_250_000,
        category_code="oylik",
        worker="Akmal",
        pay_type="Avans",
        currency="USD",
        orig_amount=100,
        rate=12500,
        on_date=date(2026, 11, 1),
        actor_admin=True,
    )
    upload(admin_api, worker, office_template())
    r = admin_api.get(f"/api/v1/finance/workers/{worker.pk}/salary/excel/?month=2026-11")
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb["Oylik"]
    assert ws["C2"].value == "Ofis"  # filled by hand in the template: kept
    assert (ws["C3"].value, ws["C4"].value) == ("Akmal", "Usta")
    assert ws["G3"].value.date() == date(2026, 10, 1)  # pay day 5 Nov = October's pay
    assert ws["G4"].value.date() == date(2026, 11, 5)
    assert wb["Sozlamalar"]["B2"].value == 12500
    # 3 advances, 2 rows: the last row carries the rest, the total stays right
    assert ws["B24"].value.date() == date(2026, 10, 6)
    assert (ws["E24"].value, ws["F24"].value, ws["G24"].value) == ("UZS", 300_000, None)
    assert (ws["E25"].value, ws["F25"].value) == ("UZS", 1_450_000)
    assert "+2 ta avans" in ws["C25"].value
    assert ws["H24"].value.startswith("=IF(") and ws["A26"].value == "JAMI AVANSLAR"


def test_office_template_usd_row(admin_api, members, monkeypatch):
    monkeypatch.setattr(S, "usd_rate", lambda: S.D(12500))
    w = Worker.objects.create(group="Ofis", name="Dilnoza", pay_day=10)
    S.create_operation(
        member=members[0],
        kind="expense",
        amount=1_250_000,
        category_code="oylik",
        worker="Dilnoza",
        pay_type="Avans",
        currency="USD",
        orig_amount=100,
        rate=12500,
        on_date=date(2026, 11, 1),
        actor_admin=True,
    )
    upload(admin_api, w, office_template())
    r = admin_api.get(f"/api/v1/finance/workers/{w.pk}/salary/excel/?month=2026-11")
    ws = load_workbook(io.BytesIO(r.content))["Oylik"]
    assert (ws["E24"].value, ws["F24"].value, ws["G24"].value) == ("USD", 100, 12500)
    assert ws["B25"].value is None


def test_own_period_and_untyped_payments(admin_api, worker, members):
    old = S.create_operation(
        member=members[0],
        kind="expense",
        amount=70_000,
        category_code="oylik",
        worker="Akmal",
        on_date=date(2026, 10, 10),
        actor_admin=True,
    )
    type(old).objects.filter(pk=old.pk).update(pay_type="")  # an imported row
    base = f"/api/v1/finance/workers/{worker.pk}/salary/"
    r = admin_api.get(base + "?month=2026-11&date_from=2026-10-01&date_to=2026-10-31")
    assert (r.data["date_from"], r.data["date_to"]) == (date(2026, 10, 1), date(2026, 10, 31))
    assert r.data["advance_total"] == 501_000  # 5, 6 and 20 Oct
    assert [x["amount"] for x in r.data["untyped"]] == [70_000]
    assert admin_api.get(base + "?date_from=2026-10-01").status_code == 400
    r = admin_api.get(
        f"/api/v1/finance/workers/{worker.pk}/salary/excel/"
        "?month=2026-11&date_from=2026-10-01&date_to=2026-10-31"
    )
    assert load_workbook(io.BytesIO(r.content))["Avanslar"]["C7"].value == 501_000

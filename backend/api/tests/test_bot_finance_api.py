"""Finance part of the internal bot API (/api/v1/bot/finance/...): the rules the Telegram
router used to hold, now enforced by the backend. Members: 100 boss (admin), 200 Ali
(staff), 300 Vali (staff, bonus receiver)."""

from datetime import date, timedelta
from decimal import Decimal

import pytest
from constance import config
from django.core.files.uploadedfile import SimpleUploadedFile

from accounting import followup
from accounting import services as S
from accounting.models import (
    ClientFeedback,
    ManualOrder,
    MoneyRequest,
    MonthReport,
    Operation,
    OrderBonus,
    SiteEntryTask,
)
from oauth.models import TelegramMember
from order.models import Order

pytestmark = pytest.mark.django_db

BOSS, ALI, VALI = 100, 200, 300


@pytest.fixture(autouse=True)
def _config(db):
    config.FIN_USD_RATE = 12000.0
    config.FIN_CHANNEL_ID = ""


def op_body(**kw):
    return {"kind": "expense", "amount": "1000", "category": "ovqat", **kw}


def book(bot, actor=ALI, **kw):
    return bot.post("finance/operations/", actor=actor, data=op_body(**kw))


# ================================================================== whoami / join / roles
def test_whoami_refreshes_username_and_strangers_get_null(bot_client, members):
    r = bot_client.post("whoami/", actor=ALI, data={"user": {"id": ALI, "username": "ali_u"}})
    assert r.status_code == 200
    assert r.data["finance_member"]["username"] == "ali_u"
    assert r.data["finance_member"]["telegram_id"] == ALI and r.data["hr"]["is_admin"] is False
    assert TelegramMember.objects.get(telegram_id=ALI).username == "ali_u"
    boss = bot_client.post("whoami/", actor=BOSS, data={"user": {"id": BOSS}})
    assert boss.data["hr"]["is_admin"] is True
    stranger = bot_client.post("whoami/", actor=777, data={"user": {"id": 777}})
    assert stranger.data == {"finance_member": None, "hr": {"employee": None, "is_admin": False}}
    assert bot_client.post("whoami/", actor=ALI, data={"user": {"id": 999}}).status_code == 403


def test_join_creates_pending_member_and_asks_the_admins(bot_client, members, bot_events):
    r = bot_client.post("finance/join/", actor=555, data={"name": "Sardor", "username": "sardor"})
    assert r.status_code == 201 and r.data["created"] is True
    m = TelegramMember.objects.get(telegram_id=555)
    assert m.role == "pending" and m.username == "sardor"
    ev = bot_events.one("finance.join_requested")
    assert ev["admins"] == [BOSS] and ev["superusers"] == []
    assert ev["member"]["telegram_id"] == 555
    again = bot_client.post("finance/join/", actor=555, data={"name": "Sardor"})
    assert again.data["created"] is False and len(bot_events.of("finance.join_requested")) == 1


def test_join_without_admins_asks_the_superusers_and_never_links_a_site_account(
    bot_client, members, bot_events, make_staff
):
    TelegramMember.objects.filter(role="admin").update(role="staff")
    owner = make_staff(superuser=True, telegram_id=555)  # linked by an unverified /start <id>
    bot_client.post("finance/join/", actor=555, data={"name": "Sardor"})
    assert bot_events.one("finance.join_requested")["superusers"] == [555]
    assert TelegramMember.objects.get(telegram_id=555).user_id is None
    assert owner.telegram_id == 555


def test_admin_approves_and_blocks(bot_client, members, bot_events):
    TelegramMember.objects.create(telegram_id=555, name="Sardor", role="pending")
    assert (
        bot_client.post("finance/members/555/role/", actor=ALI, data={"role": "staff"}).status_code
        == 403
    )
    assert (
        bot_client.post("finance/members/555/role/", actor=BOSS, data={"role": "admin"}).status_code
        == 400
    )
    r = bot_client.post("finance/members/555/role/", actor=BOSS, data={"role": "staff"})
    assert r.status_code == 200 and r.data["role"] == "staff"
    assert bot_events.one("finance.member_approved")["member"]["telegram_id"] == 555
    r = bot_client.post("finance/members/555/role/", actor=BOSS, data={"role": "blocked"})
    assert r.data["role"] == "blocked" and len(bot_events.of("finance.member_approved")) == 1


def test_duty_is_exclusive_admin_only_and_task_only_cannot_book(bot_client, members, bot_events):
    assert (
        bot_client.post(
            "finance/members/200/duty/", actor=ALI, data={"duty": "is_site_entry"}
        ).status_code
        == 403
    )
    assert (
        bot_client.post(
            "finance/members/200/duty/", actor=BOSS, data={"duty": "is_bonus_receiver"}
        ).status_code
        == 400
    )
    r = bot_client.post("finance/members/200/duty/", actor=BOSS, data={"duty": "is_site_entry"})
    assert r.status_code == 200 and r.data["is_task_only"] and not r.data["has_wallet"]
    assert bot_events.one("finance.duty_assigned")["duty"] == "is_site_entry"
    assert book(bot_client).status_code == 403  # task-only members don't book money
    bot_client.post("finance/members/300/duty/", actor=BOSS, data={"duty": "is_site_entry"})
    assert not TelegramMember.objects.get(telegram_id=ALI).is_site_entry  # one person per duty


def test_rename_member_admin_only(bot_client, members):
    assert (
        bot_client.post("finance/members/200/name/", actor=ALI, data={"name": "X"}).status_code
        == 403
    )
    r = bot_client.post("finance/members/200/name/", actor=BOSS, data={"name": "Jasur (zakupshik)"})
    assert r.data["name"] == "Jasur (zakupshik)"
    assert (
        bot_client.post("finance/members/200/name/", actor=BOSS, data={"name": " "}).status_code
        == 400
    )


def test_members_lists_follow_roles(bot_client, members):
    TelegramMember.objects.create(telegram_id=555, name="Sardor", role="pending")
    pending = bot_client.get("finance/members/?set=pending", actor=BOSS)
    assert [m["telegram_id"] for m in pending.data] == [555]
    assert bot_client.get("finance/members/?set=pending", actor=ALI).status_code == 403
    wallets = bot_client.get("finance/members/?set=wallets", actor=ALI)
    assert {m["telegram_id"] for m in wallets.data} == {BOSS, ALI, VALI}
    staff = bot_client.get("finance/members/?set=staff", actor=BOSS)
    assert {m["telegram_id"] for m in staff.data} == {ALI, VALI}
    duty = bot_client.get("finance/members/?duty=is_bonus_receiver", actor=BOSS)
    assert [m["telegram_id"] for m in duty.data] == [VALI]
    assert bot_client.get("finance/members/200/", actor=ALI).status_code == 200
    assert bot_client.get("finance/members/300/", actor=ALI).status_code == 403
    assert bot_client.get("finance/members/300/balance/", actor=ALI).status_code == 403
    assert bot_client.get("finance/members/200/balance/", actor=ALI).data == {"balance": "0.00"}


# ================================================================== catalogs, parse, config
def test_catalogs_config_and_parse(bot_client, members):
    cats = bot_client.get("finance/categories/?kind=expense", actor=ALI).data
    assert cats and all(c["kind"] == "expense" for c in cats)
    assert {"code", "name", "emoji", "admin_only", "no_order", "archive", "excluded"} <= set(
        cats[0]
    )
    assert bot_client.get("finance/workers/", actor=ALI).data
    assert bot_client.get("finance/cars/", actor=ALI).data
    cfg = bot_client.get("finance/config/", actor=ALI).data
    assert cfg["usd_rate"] == "12000.0000" and cfg["usd_rate_manual"] == "12000.0000"
    q = bot_client.post(
        "finance/parse/", actor=ALI, data={"text": "350000 furnitura 406 petlya"}
    ).data
    assert (q["kind"], q["cat"], q["cur"], q["val"]) == ("expense", "furnitura", "UZS", "350000")
    assert bot_client.post("finance/parse/", actor=ALI, data={"text": "salom"}).data is None
    m = bot_client.post("finance/parse/", actor=ALI, data={"text": "100$", "mode": "money"}).data
    assert m == {"amount": "1200000.00", "cur": "USD", "orig": "100.00", "rate": "12000.0000"}
    assert (
        bot_client.post("finance/parse/", actor=ALI, data={"text": "abc", "mode": "money"}).data
        is None
    )
    config.FIN_USD_RATE = 0.0
    config.USD_CURRENCY = 0
    from django.core.cache import cache

    cache.clear()
    cache.set("fin_cbu_down", "1", 60)
    r = bot_client.post("finance/parse/", actor=ALI, data={"text": "100$", "mode": "money"})
    assert r.status_code == 400  # no rate for dollars


def test_usd_rate_is_admin_only(bot_client, members):
    assert (
        bot_client.put("finance/config/usd-rate/", actor=ALI, data={"rate": "12500"}).status_code
        == 403
    )
    r = bot_client.put("finance/config/usd-rate/", actor=BOSS, data={"rate": "12500"})
    assert r.status_code == 200 and r.data["usd_rate_manual"] == "12500.0000"
    assert (
        bot_client.put("finance/config/usd-rate/", actor=BOSS, data={"rate": "abc"}).status_code
        == 400
    )
    bot_client.put("finance/config/usd-rate/", actor=BOSS, data={"rate": "0"})
    assert float(config.FIN_USD_RATE) == 0.0


def test_report_template_is_admin_only(bot_client, members, settings, tmp_path):
    settings.MEDIA_ROOT = str(tmp_path)
    png = SimpleUploadedFile("t.png", b"\x89PNG fake", content_type="image/png")
    assert (
        bot_client.call(
            "POST", "finance/report-template/", actor=ALI, files={"file": png}
        ).status_code
        == 403
    )
    png = SimpleUploadedFile("t.png", b"\x89PNG fake", content_type="image/png")
    assert (
        bot_client.call(
            "POST", "finance/report-template/", actor=BOSS, files={"file": png}
        ).status_code
        == 204
    )
    assert config.FIN_REPORT_TEMPLATE
    assert bot_client.delete("finance/report-template/", actor=ALI).status_code == 403
    assert bot_client.delete("finance/report-template/", actor=BOSS).status_code == 204
    assert config.FIN_REPORT_TEMPLATE == ""


# ================================================================== operations
def test_quick_expense_on_site_order_emits_one_event(bot_client, members, site_order, bot_events):
    r = book(
        bot_client,
        kind="expense",
        amount="350000",
        category="furnitura",
        order_ref="406",
        comment="petlya",
    )
    assert r.status_code == 201, r.data
    op = Operation.objects.get()
    assert (op.kind, op.amount, op.category.code, op.order_ref, op.order_id) == (
        "expense",
        Decimal("350000"),
        "furnitura",
        "406",
        site_order.id,
    )
    assert op.member.telegram_id == ALI and op.author.telegram_id == ALI and op.source == "bot"
    assert r.data["balance"] == "-350000.00" and r.data["operation"]["number"] == op.number
    ev = bot_events.one("finance.operation_created")
    assert (
        ev["balance"] == "-350000.00" and ev["notify_owner"] is False and ev["from_site"] is False
    )
    assert ev["owner_tg"] == ALI and ev["operation"]["category"]["code"] == "furnitura"


def test_event_is_queued_once_and_only_after_commit(
    bot_client, members, monkeypatch, django_capture_on_commit_callbacks
):
    sent = []
    monkeypatch.setattr("core.tasks.notify_bot.delay", lambda env: sent.append(env))
    with django_capture_on_commit_callbacks(execute=False) as callbacks:
        assert book(bot_client).status_code == 201
    assert sent == []  # nothing before the transaction commits
    for cb in callbacks:
        cb()
    assert [e["event"] for e in sent] == ["finance.operation_created"]
    assert sent[0]["id"] and sent[0]["data"]["operation"]["amount"] == "1000.00"


def test_dollar_expense_keeps_original_amount_and_rate(bot_client, members, site_order, bot_events):
    r = book(
        bot_client,
        amount="1200000",
        currency="USD",
        orig_amount="100",
        rate="12000",
        category="material",
        order_ref="406",
        comment="LDSP 3 list",
    )
    assert r.status_code == 201
    op = Operation.objects.get()
    assert op.currency == "USD" and op.orig_amount == 100 and op.amount == Decimal("1200000")
    assert book(bot_client, currency="USD", amount="1").status_code == 400  # dollars need the rate


def test_salary_keeps_worker_and_pay_type_and_never_an_order(bot_client, members, bot_events):
    r = book(
        bot_client,
        actor=BOSS,
        amount="2000000",
        category="oylik",
        worker="Dilmurod",
        pay_type="Oylik",
        order_ref="406",
        comment="oktabr",
    )
    assert r.status_code == 201
    op = Operation.objects.get()
    assert op.worker == "Dilmurod" and op.pay_type == "Oylik" and op.order_ref == ""
    staff = book(bot_client, amount="500000", category="oylik", worker="Dilmurod", pay_type="Oylik")
    assert staff.data["operation"]["pay_type"] == "Avans"  # only admins choose avans / oylik


def test_staff_cannot_book_admin_only_categories_and_cannot_spoof_admin(bot_client, members):
    assert book(bot_client, category="uy").status_code == 400
    assert book(bot_client, category="uy", actor_admin=True, internal=True).status_code == 400
    assert book(bot_client, kind="income", category="rahbar").status_code == 400
    assert book(bot_client, category="otkazma").status_code == 400  # only via transfers
    old = str(date.today() - timedelta(days=10))
    assert (
        book(bot_client, date=old).status_code == 201
        and Operation.objects.get().date == date.today()
    )
    assert Operation.objects.filter(category__code__in=["uy", "rahbar", "otkazma"]).count() == 0


def test_staff_cannot_book_into_somebody_elses_wallet(bot_client, members):
    r = bot_client.post("finance/operations/", actor=ALI, data=op_body(member_tg=VALI))
    assert r.status_code == 403 and not Operation.objects.exists()


def test_boss_gives_money_and_the_owner_is_told(bot_client, members, bot_events):
    r = bot_client.post(
        "finance/operations/",
        actor=BOSS,
        data={"member_tg": ALI, "kind": "income", "amount": "1000000", "category": "rahbar"},
    )
    assert r.status_code == 201
    op = Operation.objects.get(category__code="rahbar")
    assert (
        op.member.telegram_id == ALI
        and op.author.telegram_id == BOSS
        and op.category.code == "rahbar"
    )
    ev = bot_events.one("finance.operation_created")
    assert ev["notify_owner"] is True and ev["owner_tg"] == ALI
    # the money left the boss's own wallet: a paired chiqim, not a company expense
    out = op.pair
    assert out.member.telegram_id == BOSS and out.kind == "expense" and out.amount == op.amount
    assert out.category.code == "otkazma"
    assert S.balance(members_by(ALI)) == 1000000 and S.balance(members_by(BOSS)) == -1000000
    assert S.summary(None, None)["expense"] == 0
    # deleting the kirim takes the boss's chiqim with it
    S.delete_operation(op)
    assert S.balance(members_by(BOSS)) == 0 and S.balance(members_by(ALI)) == 0


def test_boss_giving_to_himself_books_no_pair(bot_client, members, bot_events):
    r = bot_client.post(
        "finance/operations/",
        actor=BOSS,
        data={"member_tg": BOSS, "kind": "income", "amount": "5000", "category": "rahbar"},
    )
    assert r.status_code == 201 and Operation.objects.count() == 1


def test_pending_and_unknown_members_cannot_book(bot_client, members):
    TelegramMember.objects.create(telegram_id=555, name="P", role="pending")
    assert book(bot_client, actor=555).status_code == 403
    assert book(bot_client, actor=9999).status_code == 403


def test_client_payment_rechecks_the_bonus(
    bot_client, members, site_order, bot_events, monkeypatch
):
    calls = []
    monkeypatch.setattr(followup, "bonus_check", lambda ref, ok_hour: calls.append((ref, ok_hour)))
    book(bot_client, kind="income", category="avans", order_ref="406", amount="100000")
    book(bot_client, kind="income", category="qoldiq", order_ref="406", amount="100000")
    book(bot_client, kind="expense", category="material", order_ref="406", amount="100000")
    assert calls == [("406", False), ("406", False)]


def test_delete_rules(bot_client, members, bot_events):
    ali, vali = members[1], members[2]
    mine = S.create_operation(
        member=ali, author=ali, kind="expense", amount=1000, category_code="boshqa"
    )
    old = S.create_operation(
        member=ali,
        author=ali,
        kind="expense",
        amount=1000,
        category_code="boshqa",
        on_date=date.today() - timedelta(days=5),
        actor_admin=True,
    )
    theirs = S.create_operation(
        member=vali, author=vali, kind="expense", amount=1000, category_code="boshqa"
    )
    for op in (old, theirs):
        r = bot_client.delete(f"finance/operations/{op.number}/", actor=ALI)
        assert r.status_code == 403, r.data
    assert bot_client.delete("finance/operations/99999/", actor=ALI).status_code == 400
    r = bot_client.delete(f"finance/operations/{mine.number}/", actor=ALI)
    assert r.status_code == 200 and len(r.data["deleted"]) == 1
    mine.refresh_from_db()
    assert mine.is_deleted
    ev = bot_events.one("finance.operation_deleted")
    assert ev["by_name"] == "Ali" and [o["number"] for o in ev["operations"]] == [mine.number]
    assert (
        bot_client.delete(f"finance/operations/{old.number}/", actor=BOSS).status_code == 200
    )  # admins: any


def test_operation_detail_reports_what_the_actor_may_do(bot_client, members):
    ali, boss = members[1], members[0]
    mine = S.create_operation(
        member=ali, author=ali, kind="expense", amount=1000, category_code="boshqa"
    )
    r = bot_client.get(f"finance/operations/{mine.number}/", actor=ALI)
    assert (
        r.data["can_delete"] is True
        and r.data["can_attach"] is True
        and r.data["number"] == mine.number
    )
    hers = S.create_operation(
        member=boss,
        author=boss,
        kind="expense",
        amount=1000,
        category_code="boshqa",
        actor_admin=True,
    )
    assert bot_client.get(f"finance/operations/{hers.number}/", actor=ALI).status_code == 403
    assert (
        bot_client.get(f"finance/operations/{hers.number}/", actor=BOSS).data["can_attach"] is True
    )


def test_receipt_only_on_own_entry(bot_client, members, bot_events):
    boss = members[0]
    op = S.create_operation(
        member=boss,
        author=boss,
        kind="expense",
        amount=1000,
        category_code="ovqat",
        actor_admin=True,
    )
    r = bot_client.post(
        f"finance/operations/{op.number}/receipt/", actor=ALI, data={"file_id": "F1"}
    )
    assert r.status_code == 403
    mine = book(bot_client).data["operation"]["number"]
    r = bot_client.post(f"finance/operations/{mine}/receipt/", actor=ALI, data={"file_id": "F1"})
    assert r.status_code == 200 and Operation.objects.get(number=mine).receipt_file_id == "F1"
    ev = bot_events.one("finance.receipt_attached")
    assert ev["file_id"] == "F1" and ev["operation"]["number"] == mine
    assert (
        bot_client.post(
            f"finance/operations/{op.number}/receipt/", actor=BOSS, data={"file_id": "F2"}
        ).status_code
        == 200
    )
    assert (
        bot_client.post(f"finance/operations/{mine}/receipt/", actor=ALI, data={}).status_code
        == 400
    )


def test_receipt_photo_with_caption_flow(bot_client, members):
    r = book(
        bot_client, amount="120000", category="ovqat", comment="tushlik", receipt_file_id="FILE9"
    )
    op = Operation.objects.get(number=r.data["operation"]["number"])
    assert op.receipt_file_id == "FILE9" and op.category.code == "ovqat"


# ================================================================== transfers
def test_transfer_writes_both_sides_and_tells_admins_once(bot_client, members, bot_events):
    r = bot_client.post(
        "finance/transfers/",
        actor=ALI,
        data={"to_tg": VALI, "amount": "300000", "comment": "zapravkaga"},
    )
    assert r.status_code == 201 and r.data["balance"] == "-300000.00"
    out, inc = Operation.objects.order_by("number")
    assert (out.kind, inc.kind, out.category.code) == ("expense", "income", "otkazma")
    assert inc.pair_id == out.id and S.balance(members[2]) == 300000
    ev = bot_events.one("finance.transfer_created")
    assert (
        ev["admins"] == [BOSS]
        and ev["inc_balance"] == "300000.00"
        and ev["inc_member"]["telegram_id"] == VALI
    )
    # deleting one half deletes both and announces once
    d = bot_client.delete(f"finance/operations/{out.number}/", actor=ALI)
    assert len(d.data["deleted"]) == 2
    inc.refresh_from_db()
    assert inc.is_deleted
    assert len(bot_events.of("finance.operation_deleted")) == 1


def test_transfer_to_blocked_or_self_is_refused(bot_client, members):
    TelegramMember.objects.create(telegram_id=901, name="Blok", role="blocked")
    assert (
        bot_client.post(
            "finance/transfers/", actor=ALI, data={"to_tg": 901, "amount": "1000"}
        ).status_code
        == 400
    )
    assert (
        bot_client.post(
            "finance/transfers/", actor=ALI, data={"to_tg": ALI, "amount": "1000"}
        ).status_code
        == 400
    )
    assert (
        bot_client.post(
            "finance/transfers/", actor=ALI, data={"to_tg": 12345, "amount": "1000"}
        ).status_code
        == 404
    )
    assert not Operation.objects.exists()


# ================================================================== orders, reports
def test_order_card_hides_contract_money_from_staff(bot_client, members, site_order):
    S.create_operation(
        member=members[1], kind="expense", amount=300000, category_code="material", order_ref="406"
    )
    staff = bot_client.get("finance/orders/406/", actor=ALI).data
    for hidden in ("total", "paid", "rest", "price_uzs", "currency"):
        assert hidden not in staff
    assert (
        staff["stats"]["spent"] == "300000.00" and staff["stats"]["cats"][0]["code"] == "material"
    )
    assert staff["stats"]["last_ops"][0]["comment"] == ""
    admin = bot_client.get("finance/orders/406/", actor=BOSS).data
    assert (admin["total"], admin["paid"], admin["rest"], admin["currency"]) == (
        "1000.00",
        "400.00",
        "600.00",
        "USD",
    )
    assert bot_client.get("finance/orders/9999/", actor=ALI).status_code == 404


def test_order_lists(bot_client, members, site_order):
    got = bot_client.get("finance/orders/?buckets=jarayonda,yangi", actor=ALI).data
    assert [o["ref"] for o in got] == ["406"] and got[0]["bucket"] == "jarayonda"
    assert bot_client.get("finance/orders/?buckets=nope", actor=ALI).status_code == 400
    assert bot_client.get("finance/orders/archive/", actor=ALI).data == []
    assert bot_client.get("finance/orders/next-ref/", actor=ALI).status_code == 403
    assert bot_client.get("finance/orders/next-ref/", actor=BOSS).data == {"ref": "SM-001"}
    assert bot_client.get("finance/orders/closed-manual/", actor=ALI).status_code == 403


def test_manual_orders_admin_only(bot_client, members, bot_events):
    body = {"ref": "SM-001", "client": "Nazira", "phone": "", "price_text": "8 mln"}
    assert bot_client.post("finance/manual-orders/", actor=ALI, data=body).status_code == 403
    r = bot_client.post("finance/manual-orders/", actor=BOSS, data=body)
    assert r.status_code == 201 and r.data["ref"] == "SM-001"
    assert ManualOrder.objects.get().price == 8_000_000
    assert bot_events.one("finance.manual_order_created")["order"]["price"] == "8000000.00"
    assert (
        bot_client.post("finance/manual-orders/", actor=BOSS, data=body).status_code == 400
    )  # id taken
    r = bot_client.patch("finance/manual-orders/SM-001/", actor=BOSS, data={"price_text": "9 mln"})
    assert r.data["manual"]["price"] == "9000000.00" and r.data["total"] == "9000000.00"
    r = bot_client.patch("finance/manual-orders/SM-001/", actor=BOSS, data={"toggle_closed": True})
    assert r.data["manual"]["is_closed"] is True
    assert bot_events.one("finance.manual_order_closed")["order"]["is_closed"] is True
    assert [o["ref"] for o in bot_client.get("finance/orders/closed-manual/", actor=BOSS).data] == [
        "SM-001"
    ]
    r = bot_client.patch("finance/manual-orders/SM-001/", actor=BOSS, data={"toggle_closed": True})
    assert r.data["manual"]["is_closed"] is False


def test_summary_and_operations_never_leak_other_wallets(bot_client, members):
    boss, ali, vali = members
    S.create_operation(member=ali, kind="expense", amount=1000, category_code="ovqat")
    S.create_operation(member=vali, kind="expense", amount=2000, category_code="ovqat")
    month = f"{date.today():%Y-%m}"
    mine = bot_client.get(f"finance/operations/?month={month}&member=all", actor=ALI).data
    assert [r["member"]["telegram_id"] for r in mine["rows"]] == [ALI] and mine["scope"] == "200"
    everyone = bot_client.get(f"finance/operations/?month={month}&member=all", actor=BOSS).data
    assert {r["member"]["telegram_id"] for r in everyone["rows"]} == {ALI, VALI}
    assert everyone["rows"][0]["signed_amount"].startswith("-")
    only_vali = bot_client.get(f"finance/operations/?month={month}&member={VALI}", actor=BOSS).data
    assert len(only_vali["rows"]) == 1
    sm = bot_client.get(f"finance/summary/?month={month}&member=all", actor=ALI).data
    assert (
        sm["scope"] == "200"
        and sm["summary"]["expense"] == "1000.00"
        and sm["member_balance"] == "-1000.00"
    )
    sm = bot_client.get(f"finance/summary/?month={month}&member=all", actor=BOSS).data
    assert (
        sm["scope"] == "all"
        and sm["summary"]["expense"] == "3000.00"
        and sm["member_balance"] is None
    )
    assert bot_client.get("finance/summary/?month=2026-13", actor=BOSS).status_code == 400
    months = bot_client.get(f"finance/months/?member={VALI}", actor=ALI).data
    assert months["scope"] == "200" and len(months["series"]) == 7


def test_rahbar_income_is_not_company_income(members):
    boss, ali, _ = members
    S.create_operation(
        member=ali,
        author=boss,
        kind="income",
        amount=1_000_000,
        category_code="rahbar",
        actor_admin=True,
    )
    S.create_operation(
        member=ali, author=ali, kind="expense", amount=200_000, category_code="ovqat"
    )
    S.create_operation(
        member=boss,
        author=boss,
        kind="expense",
        amount=50_000,
        category_code="uy",
        actor_admin=True,
    )
    sm = S.summary(None, None)
    assert (
        sm["income"] == 0
        and sm["expense"] == 200000
        and sm["given_to_staff"] == 1000000
        and sm["home"] == 50000
    )


def test_wallets_checks_excel_and_handover_list_are_admin_only(bot_client, members, site_order):
    S.create_operation(member=members[1], kind="expense", amount=5000, category_code="ovqat")
    for path in (
        "finance/wallets/",
        "finance/checks/",
        "finance/handover/",
        f"finance/excel/?month={date.today():%Y-%m}",
    ):
        assert bot_client.get(path, actor=ALI).status_code == 403, path
    w = bot_client.get("finance/wallets/", actor=BOSS).data
    rows = {r["member"]["name"]: r for r in w["rows"]}
    assert rows["Ali"]["balance"] == "-5000.00" and rows["Ali"]["month_spent"] == "5000.00"
    assert w["pending"] == [] and w["staff_total"] == "-5000.00" and "report_month" in w
    assert isinstance(bot_client.get("finance/checks/", actor=BOSS).data, list)
    x = bot_client.get(f"finance/excel/?month={date.today():%Y-%m}", actor=BOSS)
    assert x.status_code == 200 and x.content[:2] == b"PK"
    assert "attachment; filename=" in x["Content-Disposition"] and x[
        "Content-Disposition"
    ].endswith('.xlsx"')


def test_month_report_submit(bot_client, members, bot_events):
    S.create_operation(member=members[1], kind="income", amount=1000, category_code="boshqa_kirim")
    g = bot_client.get("finance/month-report/", actor=ALI).data
    assert g["income"] == "1000.00" and g["submitted"] is False
    r = bot_client.post(
        "finance/month-report/", actor=ALI, data={"month": g["month"], "note": "hammasi"}
    )
    assert r.status_code == 200
    rep = MonthReport.objects.get()
    assert rep.member.telegram_id == ALI and rep.balance == 1000 and rep.note == "hammasi"
    ev = bot_events.one("finance.month_submitted")
    assert ev["closing"] == "1000.00" and ev["admins"] == [BOSS] and ev["summary"]["count"] == 1
    assert bot_client.get("finance/month-report/", actor=ALI).data["submitted"] is True
    assert (
        bot_client.post("finance/month-report/", actor=ALI, data={"month": "2999-01"}).status_code
        == 400
    )


# ================================================================== money requests
def test_money_request_given_by_boss(bot_client, members, bot_events):
    r = bot_client.post(
        "finance/requests/", actor=ALI, data={"amount_text": "500 ming", "reason": "406 ga LDSP"}
    )
    assert (
        r.status_code == 201 and r.data["amount"] == "500000.00" and r.data["status"] == "pending"
    )
    num = r.data["number"]
    assert bot_events.one("finance.request_created")["admins"] == [BOSS]
    assert bot_client.post(f"finance/requests/{num}/give/", actor=ALI).status_code == 403
    assert bot_client.get(f"finance/requests/{num}/", actor=ALI).status_code == 403
    msgs = bot_client.post(
        f"finance/requests/{num}/admin-messages/", actor=ALI, data={"messages": [[BOSS, 9001]]}
    )
    assert msgs.status_code == 204 and MoneyRequest.objects.get().admin_messages == [[BOSS, 9001]]
    assert (
        bot_client.post(
            f"finance/requests/{num}/admin-messages/", actor=VALI, data={"messages": []}
        ).status_code
        == 403
    )
    g = bot_client.post(f"finance/requests/{num}/give/", actor=BOSS, data={})
    assert g.status_code == 200 and g.data["balance"] == "500000.00"
    req = MoneyRequest.objects.get()
    assert req.status == "given" and req.operation.category.code == "rahbar"
    ev = bot_events.one("finance.request_answered")
    assert ev["status"] == "given" and ev["post_to_channel"] is True and ev["from_site"] is False
    assert (
        ev["request"]["admin_messages"] == [[BOSS, 9001]]
        and ev["operation"]["number"] == req.operation.number
    )
    # answered once: a second give / reject is 409
    assert bot_client.post(f"finance/requests/{num}/give/", actor=BOSS).status_code == 409
    assert bot_client.post(f"finance/requests/{num}/reject/", actor=BOSS).status_code == 409
    assert S.balance(members[1]) == 500000


def test_money_request_custom_amount_and_reject(bot_client, members, bot_events):
    num = bot_client.post("finance/requests/", actor=ALI, data={"amount_text": "300$"}).data[
        "number"
    ]
    g = bot_client.post(f"finance/requests/{num}/give/", actor=BOSS, data={"amount_text": "250$"})
    assert g.status_code == 200 and S.balance(members[1]) == 250 * 12000
    num2 = bot_client.post("finance/requests/", actor=ALI, data={"amount_text": "100000"}).data[
        "number"
    ]
    r = bot_client.post(f"finance/requests/{num2}/reject/", actor=BOSS)
    assert r.status_code == 200 and r.data["status"] == "rejected"
    rejected = [e for e in bot_events.of("finance.request_answered") if e["status"] == "rejected"]
    assert len(rejected) == 1 and rejected[0]["post_to_channel"] is False
    assert (
        bot_client.post("finance/requests/", actor=ALI, data={"amount_text": "abc"}).status_code
        == 400
    )


# ================================================================== Готов, bonus, «Pul oldim»
@pytest.fixture
def ready(site_order, members, bot_events):
    """Order 406 became Готов: the hourly pass opens bonus tracking and the card event."""
    config.FIN_BONUS_SINCE = (date.today() - timedelta(days=1)).isoformat()
    site_order.status = "done"
    site_order.save()
    followup.hourly_orders(ok_hour=True)
    return site_order


def test_ready_order_opens_bonus_and_sends_one_card_event(ready, bot_events):
    bn = OrderBonus.objects.get(order_ref="406")
    assert bn.total == 1000 and bn.currency == "USD" and bn.status == "waiting"
    ev = bot_events.one("finance.ready_card")
    assert ev["to"] == [VALI] and ev["fallback"] is False and ev["bonus"]["order_ref"] == "406"
    assert ev["money_state"] == {"paid": "400.00", "rest": "600.00", "src": "sayt"}
    followup.hourly_orders(ok_hour=True)
    assert len(bot_events.of("finance.ready_card")) == 1  # once per order


def test_old_ready_orders_get_no_card(site_order, members, bot_events):
    config.FIN_BONUS_SINCE = (date.today() + timedelta(days=1)).isoformat()
    site_order.status = "done"
    site_order.save()
    followup.hourly_orders(ok_hour=True)
    assert not OrderBonus.objects.exists() and not bot_events.of("finance.ready_card")


def test_collect_full_books_money_opens_site_task_and_asks_for_handover(
    bot_client, ready, bot_events
):
    assert (
        bot_client.post(
            "finance/bonuses/406/collect/", actor=ALI, data={"amount_text": "100$"}
        ).status_code
        == 403
    )
    r = bot_client.post(
        "finance/bonuses/406/collect/", actor=VALI, data={"amount_text": "600$", "full": True}
    )
    assert r.status_code == 200, r.data
    assert (
        r.data["full"] is True and r.data["bonus"]["status"] == "done" and r.data["rest"] == "0.00"
    )
    bn = OrderBonus.objects.get(order_ref="406")
    assert bn.percent == Decimal("0.5") and bn.bonus == Decimal("60000")  # 1000$ * 12000 * 0.5%
    op = Operation.objects.get(category__code="qoldiq")
    assert op.member.telegram_id == VALI and op.orig_amount == 600
    # the site's «Полученная сумма» follows: 400$ before + 600$ collected
    site_order = Order.objects.get(order_number=406)
    assert site_order.lost_money.amount == 1000 and str(site_order.lost_money.currency) == "USD"
    task = SiteEntryTask.objects.get(order_ref="406")
    assert task.is_done is True and task.done_by == "Avtomatik" and r.data["task"]["code"] == task.code
    assert not bot_events.of("finance.site_entry_created")  # nothing to enter by hand
    ho = bot_events.one("finance.handover_requested")
    assert ho["admins"] == [BOSS] and ho["ref"] == "406" and ho["operation"]["number"] == op.number
    done = bot_events.one("finance.bonus_done")
    assert done["to"] == [VALI] and done["also_admins"] == [BOSS]
    # the boss accepts the cash; a second time is 409
    assert bot_client.post(f"finance/handover/{op.number}/accept/", actor=VALI).status_code == 403
    a = bot_client.post(f"finance/handover/{op.number}/accept/", actor=BOSS)
    assert a.status_code == 200 and a.data["client"] and a.data["member_balance"] == "0.00"
    assert S.balance(members_by(300)) == 0 and S.balance(members_by(100)) == op.amount
    assert bot_client.post(f"finance/handover/{op.number}/accept/", actor=BOSS).status_code == 409
    assert bot_client.post(f"finance/handover/{op.number}/drop/", actor=BOSS).status_code == 409
    assert bot_events.one("finance.handover_resolved")["action"] == "accept"


def members_by(tg):
    return TelegramMember.objects.get(telegram_id=tg)


def test_partial_collect_then_close_rules(bot_client, ready, bot_events):
    r = bot_client.post("finance/bonuses/406/collect/", actor=VALI, data={"amount_text": "100$"})
    assert r.status_code == 200 and r.data["full"] is False and r.data["rest"] == "500.00"
    assert "bonus_done" not in "".join(bot_events.names())
    assert (
        bot_client.post("finance/bonuses/406/close/", actor=VALI).status_code == 400
    )  # rest > tolerance
    assert bot_client.post("finance/bonuses/406/close/", actor=ALI).status_code == 403
    bot_client.post("finance/bonuses/406/collect/", actor=VALI, data={"amount_text": "500$"})
    assert OrderBonus.objects.get(order_ref="406").status == "done"
    assert bot_client.post("finance/bonuses/406/close/", actor=VALI).status_code == 409
    assert (
        bot_client.post(
            "finance/bonuses/406/collect/", actor=VALI, data={"amount_text": "1$"}
        ).status_code
        == 409
    )


def test_bonus_close_when_the_site_shows_nothing_owed(bot_client, ready, bot_events, site_order):
    site_order.lost_money = site_order.price
    site_order.save()
    r = bot_client.post("finance/bonuses/406/close/", actor=BOSS)
    assert r.status_code == 200 and r.data["status"] == "done"
    assert len(bot_events.of("finance.bonus_done")) == 1


def test_bonus_read_endpoints(bot_client, ready):
    g = bot_client.get("finance/bonuses/406/", actor=VALI).data
    assert g["money_state"]["rest"] == "600.00" and g["is_receiver"] is True
    assert g["tolerance"] == "0.50" and g["usd_rate"] == "12000.0000" and g["percent_now"] == "0.50"
    assert bot_client.get("finance/bonuses/406/", actor=ALI).status_code == 403
    assert bot_client.get("finance/bonuses/777/", actor=VALI).status_code == 404
    png = bot_client.get("finance/bonuses/406/card.png", actor=VALI)
    assert (
        png.status_code == 200
        and png["Content-Type"] == "image/png"
        and png.content[:4] == b"\x89PNG"
    )


def test_card_png_for_an_order_without_a_bonus(bot_client, members, site_order):
    assert bot_client.get("finance/bonuses/406/card.png", actor=BOSS).status_code == 404
    r = bot_client.get("finance/bonuses/406/card.png?or_temp=1", actor=BOSS)
    assert r.status_code == 200 and r.content[:4] == b"\x89PNG"
    t = bot_client.get("finance/bonuses/406/?or_temp=1", actor=BOSS).data
    assert t["bonus"]["order_ref"] == "406" and not OrderBonus.objects.exists()


def test_handover_rejects_non_collection_rows_and_unknown(bot_client, members):
    op = S.create_operation(
        member=members[1], author=members[1], kind="expense", amount=1000, category_code="ovqat"
    )
    assert bot_client.post(f"finance/handover/{op.number}/drop/", actor=BOSS).status_code == 400
    assert bot_client.post("finance/handover/9999/accept/", actor=BOSS).status_code == 404


def test_handover_drop_marks_row_and_blocks_second_answer(bot_client, members, bot_events):
    got = S.create_operation(
        member=members[1],
        author=members[1],
        kind="income",
        amount=1000,
        category_code="qoldiq",
        comment="Pul oldim (mijozdan qoldiq)",
        internal=True,
    )
    listed = bot_client.get("finance/handover/", actor=BOSS).data
    assert [o["number"] for o in listed] == [got.number]
    r = bot_client.post(f"finance/handover/{got.number}/drop/", actor=BOSS)
    assert r.status_code == 200 and r.data["operation"]["handover"] == "deleted"
    assert bot_client.post(f"finance/handover/{got.number}/accept/", actor=BOSS).status_code == 409
    assert bot_events.one("finance.handover_resolved")["action"] == "drop"


# ================================================================== warranty, tasks, feedback
def test_warranty_and_site_entry_duties(bot_client, ready, bot_events):
    assert bot_client.get("finance/warranty/?given=0", actor=ALI).status_code == 403
    rows = bot_client.get("finance/warranty/?given=0", actor=VALI).data
    assert [w["order_ref"] for w in rows] == ["406"]
    assert bot_client.post("finance/warranty/406/give/", actor=ALI).status_code == 403
    r = bot_client.post("finance/warranty/406/give/", actor=VALI)
    assert r.status_code == 200 and r.data["given"] is True
    assert bot_events.one("finance.warranty_given")["by_name"] == "Vali"
    assert bot_client.post("finance/warranty/406/give/", actor=VALI).status_code == 409
    # tasks: scoped by duty
    bot_client.post("finance/bonuses/406/collect/", actor=VALI, data={"amount_text": "100$"})
    assert SiteEntryTask.objects.get().is_done  # a site order: updated automatically
    assert Order.objects.get(order_number=406).lost_money.amount == 500
    code = S.site_entry_create("M-1", "Qo'lda", "100 $", "Vali", "0").code  # a manual order
    assert bot_client.get("finance/tasks/", actor=ALI).data == {"site_entry": [], "feedback": []}
    assert len(bot_client.get("finance/tasks/", actor=BOSS).data["site_entry"]) == 1
    assert bot_client.post(f"finance/site-entry/{code}/done/", actor=ALI).status_code == 403
    assert bot_client.post(f"finance/site-entry/{code}/done/", actor=BOSS).status_code == 200
    assert bot_events.one("finance.site_entry_done")["by_name"] == "Rahbar"
    assert bot_client.post(f"finance/site-entry/{code}/done/", actor=BOSS).status_code == 409
    assert bot_client.post("finance/site-entry/nope/done/", actor=BOSS).status_code == 404


def test_feedback_upload_creates_missing_row_and_saves_the_picture(
    bot_client, members, bot_events, settings, tmp_path
):
    settings.MEDIA_ROOT = str(tmp_path)
    TelegramMember.objects.filter(telegram_id=VALI).update(is_feedback=True)
    shot = SimpleUploadedFile("s.jpg", b"\xff\xd8\xff fake jpeg", content_type="image/jpeg")
    r = bot_client.call(
        "POST", "finance/feedback/389/", actor=ALI, files={"file": shot}, form={"file_id": "F1"}
    )
    assert r.status_code == 403  # Ali has no feedback duty
    shot = SimpleUploadedFile("s.jpg", b"\xff\xd8\xff fake jpeg", content_type="image/jpeg")
    r = bot_client.call(
        "POST",
        "finance/feedback/389/",
        actor=VALI,
        files={"file": shot},
        form={"file_id": "F1", "caption": "mijoz mamnun"},
    )
    assert r.status_code == 200, r.data
    fb = ClientFeedback.objects.get(order_ref="389")
    assert fb.status == "received" and fb.text == "mijoz mamnun" and fb.screenshot_file_id == "F1"
    assert fb.screenshot_url and (tmp_path / "finance" / "feedback").exists()
    ev = bot_events.one("finance.feedback_received")
    assert ev["file_id"] == "F1" and ev["caption"] == "mijoz mamnun" and ev["admins"] == [BOSS]
    assert bot_client.get("finance/feedback/389/", actor=VALI).data["status"] == "received"
    assert bot_client.get("finance/feedback/555/", actor=VALI).status_code == 404
    shot = SimpleUploadedFile("s.jpg", b"x", content_type="image/jpeg")
    assert (
        bot_client.call(
            "POST", "finance/feedback/389/", actor=VALI, files={"file": shot}
        ).status_code
        == 409
    )


def test_money_request_given_by_boss_leaves_his_wallet(bot_client, members, bot_events):
    bot_client.post("finance/requests/", actor=ALI, data={"amount_text": "300000", "reason": "x"})
    num = MoneyRequest.objects.get().number
    r = bot_client.post(f"finance/requests/{num}/give/", actor=BOSS, data={})
    assert r.status_code == 200, r.data
    assert S.balance(members_by(ALI)) == 300000 and S.balance(members_by(BOSS)) == -300000


def test_collect_without_site_figure_starts_from_bot_bookings(bot_client, ready, members, bot_events):
    Order.objects.filter(order_number=406).update(lost_money=None)
    S.create_operation(
        member=members_by(ALI), kind="income", amount=1200000, category_code="avans",
        currency="USD", orig_amount=100, rate=12000, order_ref="406", internal=True,
    )
    r = bot_client.post("finance/bonuses/406/collect/", actor=VALI, data={"amount_text": "200$"})
    assert r.status_code == 200, r.data
    assert Order.objects.get(order_number=406).lost_money.amount == 300  # 100$ avans + 200$


def test_collect_full_sets_the_site_paid_to_the_price(bot_client, ready, bot_events):
    r = bot_client.post(
        "finance/bonuses/406/collect/", actor=VALI, data={"amount_text": "500$", "full": True}
    )
    assert r.status_code == 200, r.data
    assert Order.objects.get(order_number=406).lost_money.amount == 1000


def test_boss_hand_out_is_not_counted_as_a_staff_transfer(bot_client, members, bot_events):
    bot_client.post(
        "finance/operations/",
        actor=BOSS,
        data={"member_tg": ALI, "kind": "income", "amount": "700000", "category": "rahbar"},
    )
    sm = S.summary(None, None)
    assert sm["given_to_staff"] == 700000 and sm["transfers"] == 0


def test_rahbar_category_cannot_be_edited_in_or_out(bot_client, members, bot_events):
    from core.exceptions import ServiceError

    bot_client.post(
        "finance/operations/",
        actor=BOSS,
        data={"member_tg": ALI, "kind": "income", "amount": "700000", "category": "rahbar"},
    )
    op = Operation.objects.get(category__code="rahbar")
    with pytest.raises(ServiceError):
        S.update_operation(op, {"category_code": "boshqa_kirim"}, actor_admin=True)
    other = S.create_operation(
        member=members_by(ALI), kind="income", amount=1000, category_code="boshqa_kirim", internal=True
    )
    with pytest.raises(ServiceError):
        S.update_operation(other, {"category_code": "rahbar"}, actor_admin=True)


def test_month_report_excel(bot_client, members):
    from io import BytesIO

    from openpyxl import load_workbook

    S.create_operation(member=members[1], kind="income", amount=1000, category_code="boshqa_kirim")
    ym = f"{date.today():%Y-%m}"
    x = bot_client.get(f"finance/month-report/excel/?month={ym}&member={ALI}", actor=BOSS)
    assert x.status_code == 200 and x["Content-Disposition"].endswith('.xlsx"')
    wb = load_workbook(BytesIO(x.content))
    assert wb.sheetnames == ["Hisobot", "Amallar"]
    assert wb["Hisobot"]["B6"].value == 1000 and wb["Hisobot"]["D6"].value == 1000
    assert wb["Amallar"]["D3"].value == 1000 and wb["Amallar"]["F3"].value == 1000
    # staff: own file only
    assert bot_client.get(f"finance/month-report/excel/?month={ym}", actor=ALI).status_code == 200
    assert (
        bot_client.get(
            f"finance/month-report/excel/?month={ym}&member={BOSS}", actor=ALI
        ).status_code
        == 403
    )

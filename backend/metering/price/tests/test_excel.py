"""Item 5: the накладной export keeps only the per-calculation "Итого" — the
per-inventory unit prices and line totals are gone, the counts stay."""

import io

import pytest
from djmoney.money import Money
from openpyxl import load_workbook

from metering.price.excel import download_inlines_excel
from metering.price.models import (
    Calculate,
    Inventory,
    InventoryInCalculate,
    InventoryType,
    ObjectType,
    Price,
)


@pytest.fixture
def price_with_one_calculation(db):
    price = Price.objects.create(metering=None)
    obj_type = ObjectType.objects.create(name="Кухня", count_name="kv")
    calc = Calculate.objects.create(
        name="Kuxnya", price=price, obj=obj_type, count=12.5, amount=Money(15000, "USD")
    )

    count_type = InventoryType.objects.create(
        name="Механизм", obj=obj_type, type=InventoryType.TYPE_COUNT
    )
    kv_type = InventoryType.objects.create(name="Тош", obj=obj_type, type=InventoryType.TYPE_KV)

    blum = Inventory.objects.create(name="BLUM", type=count_type, price=Money(7, "USD"))
    stone = Inventory.objects.create(name="Tosh xitoy", type=kv_type, price=Money(9, "USD"))

    InventoryInCalculate.objects.create(
        inventory=blum, calculate=calc, count=4, price=Money(28, "USD")
    )
    InventoryInCalculate.objects.create(
        inventory=stone, calculate=calc, count=1, price=Money(9, "USD")
    )
    return price


def _rows(response):
    workbook = load_workbook(io.BytesIO(response.content))
    return [
        [cell if cell is not None else "" for cell in row]
        for row in workbook.active.iter_rows(values_only=True)
    ]


@pytest.mark.django_db
def test_export_drops_inventory_prices_but_keeps_counts(price_with_one_calculation):
    response = download_inlines_excel(None, None, price_with_one_calculation.pk)

    rows = _rows(response)
    flat = [str(cell) for row in rows for cell in row]

    # No per-inventory unit price and no per-inventory line total anywhere.
    for money in (Money(7, "USD"), Money(28, "USD"), Money(9, "USD")):
        assert str(money) not in flat

    blum_row = next(row for row in rows if "BLUM" in row)
    assert blum_row[0] == "Механизм"
    assert blum_row[2] == 4

    stone_row = next(row for row in rows if "Tosh xitoy" in row)
    # Кв./Пог. м. inventory carries no count column at all.
    assert [cell for cell in stone_row if cell != ""] == ["Тош", "Tosh xitoy"]


@pytest.mark.django_db
def test_export_keeps_the_calculation_total(price_with_one_calculation):
    response = download_inlines_excel(None, None, price_with_one_calculation.pk)

    rows = _rows(response)

    total_row = next(row for row in rows if row and row[0] == "Итого")
    assert "15" in str(total_row[1])


@pytest.mark.django_db
def test_export_keeps_the_calculation_header_rows(price_with_one_calculation):
    response = download_inlines_excel(None, None, price_with_one_calculation.pk)

    rows = _rows(response)
    flat_first_cells = [row[0] for row in rows if row]

    assert "Kuxnya" in flat_first_cells
    assert "Тип" in flat_first_cells
    assert "kv" in flat_first_cells

from openpyxl import Workbook
from django.http import HttpResponse
from .models import Price, Calculate, InventoryInCalculate, InventoryType


def download_inlines_excel(modeladmin, request, object_id):
    price = Price.objects.get(pk=object_id)
    queryset = Calculate.objects.filter(price_id=object_id)
    wb = Workbook()
    ws = wb.active
    ws.title = str(price.metering)

    for calc in queryset:
        # Заголовки
        ws.append(
            [
                calc.name,
            ]
        )
        ws.append(["Тип", calc.obj.name])
        ws.append([calc.obj.count_name, calc.count])

        # Per-inventory prices (unit price and line total) are deliberately left out
        # of the накладной — only the per-calculation "Итого" carries money.
        inventories = (
            InventoryInCalculate.objects.filter(calculate=calc)
            .select_related("inventory__type")
            .order_by("id")
        )
        for ic in inventories:
            if ic.inventory.type.type == InventoryType.TYPE_KV:
                ws.append(
                    [
                        ic.inventory.type.name,
                        ic.inventory.name,
                    ]
                )

            elif ic.inventory.type.type == InventoryType.TYPE_COUNT:
                ws.append([ic.inventory.type.name, ic.inventory.name, ic.count])

        ws.append(["Итого", str(calc.amount)])

        ws.append([""])
        ws.append([""])

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="{price.metering}.xlsx"'

    wb.save(response)
    return response


download_inlines_excel.short_description = "📥 Скачать накладной (Excel)"


# --- The offer as a spreadsheet (REST API: GET /prices/{id}/excel/) -------------------


def build_offer_xlsx(price) -> bytes:
    """Same content as the offer PDF: room → furniture → «label: material» → price."""
    import io

    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    from .services import measure_text, offer_data, offer_lines

    data = offer_data(price)
    wb = Workbook()
    ws = wb.active
    ws.title = "Hisob-kitob"
    ws.sheet_view.showGridLines = False
    for col, width in zip("ABCD", (34, 26, 44, 16)):
        ws.column_dimensions[col].width = width

    dark, green, grey = "1F3B33", "16A34A", "F1F5F4"
    thin = Side(style="thin", color="CBD5D1")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    usd = '#,##0" $"'

    ws.append(["ТИЖОРАТ ТАКЛИФИ"])
    ws["A1"].font = Font(size=16, bold=True, color=dark)
    ws.append([f"Мижоз: {data['client']}"])
    ws.append(
        [f"Сметачи: {data['estimator'] or '—'}", "", "Сана:", data["date"].strftime("%d.%m.%Y")]
    )
    if data["variant"]:
        ws.append([f"Вариант: {data['variant']}"])
    ws.append([])

    def row(values, *, bold=False, fill=None, color=None, size=None):
        ws.append(values)
        r = ws.max_row
        for c in range(1, 5):
            cell = ws.cell(r, c)
            cell.border = box
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.font = Font(bold=bold, color=color, size=size)
            if fill:
                cell.fill = PatternFill("solid", fgColor=fill)
        ws.cell(r, 4).number_format = usd
        ws.cell(r, 4).alignment = Alignment(horizontal="right", vertical="center")
        return r

    for room in data["rooms"]:
        r = row(
            [f"{room['icon']} {room['name']}".strip(), "", "", ""],
            bold=True,
            fill=dark,
            color="FFFFFF",
            size=12,
        )
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        row(["Мебел", "Параметр", "Материал", "Нархи"], bold=True, fill=grey)
        for it in room["items"]:
            title = it["name"] + (f" ({measure_text(it)})" if measure_text(it) else "")
            lines = offer_lines(it)
            first = ws.max_row + 1
            if not lines:
                lines = [{"label": "", "value": ""}]
            for i, ln in enumerate(lines):
                row(
                    [
                        title if i == 0 else "",
                        ln["label"],
                        ln["value"],
                        it["amount"] if i == 0 else None,
                    ],
                    bold=False,
                )
            ws.cell(first, 1).font = Font(bold=True)
            ws.cell(first, 4).font = Font(bold=True)
            if len(lines) > 1:
                ws.merge_cells(start_row=first, start_column=1, end_row=ws.max_row, end_column=1)
                ws.merge_cells(start_row=first, start_column=4, end_row=ws.max_row, end_column=4)
                ws.cell(first, 1).alignment = Alignment(vertical="top", wrap_text=True)
        r = row(["Жами", "", "", room["total"]], bold=True, color=green)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        ws.append([])

    r = row(
        [
            "УМУМИЙ ЖАМИ" + (f" — {data['variant']}" if data["variant"] else ""),
            "",
            "",
            data["total"],
        ],
        bold=True,
        fill=dark,
        color="FFFFFF",
        size=13,
    )
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
